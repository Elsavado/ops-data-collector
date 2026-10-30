"""Openly licensed public data that a security / IT-services operator works with."""
import datetime as dt
import gzip
import hashlib
import io
import json
import os
import time

from docx import Document
from docx.shared import Pt

from .common import add_sheet, get, new_book, safe, save_book, session

PUB = "Public Data"
TI = f"{PUB}/1 Threat Intelligence"
STD = f"{PUB}/2 Standards & Compliance"
ITSM = f"{PUB}/3 IT Service Management"
PHYS = f"{PUB}/4 Physical Security & Surveillance"
HOSP = f"{PUB}/5 Hospitality & Maritime Operations"
CLOUD = f"{PUB}/6 Cloud & Infrastructure Operations"
BIZ = f"{PUB}/7 Workforce & Business Benchmarks"
LIB = f"{PUB}/8 Books & Knowledge Library"

OPEN_LICENSES = ("cc-by", "cc0", "cc-zero", "odc", "pddl", "mit", "apache", "bsd", "gpl", "lgpl",
                 "public", "us-pd", "ogl", "etalab", "other-open", "other-pd")

# Topic searches. Each query feeds one topic folder.
SEARCHES = {
    TI: ["intrusion detection dataset", "network traffic attack dataset", "phishing dataset",
         "malware dataset", "security operations center alerts", "cyber threat intelligence",
         "honeypot logs", "vulnerability dataset"],
    STD: ["information security policy", "security compliance audit", "risk assessment framework",
          "data breach notification"],
    ITSM: ["IT service management incident", "help desk tickets", "service desk", "IT asset inventory",
           "work orders maintenance", "service requests", "event log process mining"],
    PHYS: ["video surveillance dataset", "CCTV", "access control logs", "security guard",
           "crime incidents", "calls for service", "intrusion alarm", "fire alarm"],
    HOSP: ["hotel operations", "hotel booking", "cruise ship", "vessel inspections", "AIS vessel",
           "port operations", "tourism accommodation"],
    CLOUD: ["cloud workload trace", "server logs", "system log anomaly detection", "data center",
            "network monitoring", "uptime outage"],
    BIZ: ["security services contracts", "facility management costs", "IT spending",
          "technician wages", "occupational employment security"],
}


def meta_note(sink, folder, info):
    info = dict(info, retrieved=dt.date.today().isoformat())
    sink.write_text(f"{folder}/_source.json", json.dumps(info, indent=1, ensure_ascii=False))


# ---------- fixed sources ----------

def kev(sink, st):
    s = session()
    r = get(s, "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json")
    if not r:
        return 0
    d = r.json()
    ver = d.get("catalogVersion")
    if st.get("kev") == ver:
        return 0
    rows = [[v.get(k) for k in ("cveID", "vendorProject", "product", "vulnerabilityName", "dateAdded",
                                "shortDescription", "requiredAction", "dueDate",
                                "knownRansomwareCampaignUse", "notes")] for v in d["vulnerabilities"]]
    wb = new_book()
    add_sheet(wb, "KEV", ["CVE", "Vendor", "Product", "Name", "Date added", "Description",
                          "Required action", "Due date", "Ransomware use", "Notes"], rows,
              {"Description": 80, "Required action": 50, "Name": 40})
    by = {}
    for r_ in rows:
        by.setdefault(r_[4][:7], 0)
        by[r_[4][:7]] += 1
    add_sheet(wb, "Added per month", ["Month", "Count"], sorted(by.items()))
    base = f"{TI}/CISA Known Exploited Vulnerabilities"
    save_book(sink, f"{base}/KEV catalog {d.get('dateReleased', '')[:10]}.xlsx", wb)
    sink.write_bytes(f"{base}/json/KEV catalog {d.get('dateReleased', '')[:10]}.json.gz",
                     gzip.compress(r.content))
    meta_note(sink, base, {"source": "CISA Known Exploited Vulnerabilities Catalog",
                           "url": "https://www.cisa.gov/known-exploited-vulnerabilities-catalog",
                           "license": "US Government work, public domain"})
    st["kev"] = ver
    return len(rows)


def _cvss(m):
    for k in ("cvssMetricV40", "cvssMetricV31", "cvssMetricV30", "cvssMetricV2"):
        if m.get(k):
            c = m[k][0]["cvssData"]
            return k[10:], c.get("baseScore"), c.get("baseSeverity") or m[k][0].get("baseSeverity"), \
                c.get("vectorString")
    return None, None, None, None


def nvd(sink, st, months_per_run=12):
    """CVE records per published month. Backfills from 2015-01, then keeps the last 2 months fresh."""
    s = session()
    today = dt.date.today()
    cur = st.get("nvd_next", "2015-01")
    todo = []
    y, m = map(int, cur.split("-"))
    while (y, m) <= (today.year, today.month) and len(todo) < months_per_run:
        todo.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    last = (today.replace(day=1) - dt.timedelta(days=1))
    for ym in [(last.year, last.month), (today.year, today.month)]:
        if ym not in todo and f"{ym[0]}-{ym[1]:02d}" < cur:
            todo.append(ym)
    total = 0
    for y, m in todo:
        if sink.full():
            break
        start = dt.datetime(y, m, 1)
        end = dt.datetime(y + (m == 12), 1 if m == 12 else m + 1, 1) - dt.timedelta(milliseconds=1)
        rows, idx = [], 0
        while True:
            url = ("https://services.nvd.nist.gov/rest/json/cves/2.0?resultsPerPage=2000"
                   f"&startIndex={idx}&pubStartDate={start:%Y-%m-%dT%H:%M:%S}.000"
                   f"&pubEndDate={end:%Y-%m-%dT%H:%M:%S}.999")
            time.sleep(6.5)
            r = get(s, url, tries=6)
            if not r:
                rows = None
                break
            d = r.json()
            for v in d.get("vulnerabilities", []):
                c = v["cve"]
                desc = next((x["value"] for x in c.get("descriptions", []) if x["lang"] == "en"), "")
                ver, score, sev, vec = _cvss(c.get("metrics", {}))
                cwes = sorted({x["value"] for w in c.get("weaknesses", []) for x in w["description"]})
                cpes = set()
                for conf in c.get("configurations", []):
                    for n in conf.get("nodes", []):
                        for cm in n.get("cpeMatch", []):
                            p = cm["criteria"].split(":")
                            if len(p) > 4:
                                cpes.add(f"{p[3]}/{p[4]}")
                rows.append([c["id"], c.get("published", "")[:19], c.get("lastModified", "")[:19],
                             c.get("vulnStatus"), ver, score, sev, vec, ", ".join(cwes),
                             ", ".join(sorted(cpes)[:40]), len(c.get("references", [])), desc[:32000]])
            idx += d.get("resultsPerPage", 0)
            if idx >= d.get("totalResults", 0) or not d.get("resultsPerPage"):
                break
        if rows is None:
            break
        wb = new_book()
        add_sheet(wb, f"CVEs {y}-{m:02d}",
                  ["CVE", "Published", "Last modified", "Status", "CVSS version", "Base score", "Severity",
                   "Vector", "CWE", "Vendor/product", "References", "Description"], rows,
                  {"Description": 90, "Vendor/product": 40, "Vector": 45})
        sev = {}
        for r_ in rows:
            sev[r_[6] or "NONE"] = sev.get(r_[6] or "NONE", 0) + 1
        add_sheet(wb, "Severity summary", ["Severity", "Count"], sorted(sev.items()))
        save_book(sink, f"{TI}/NVD CVE Records/{y}/NVD CVEs {y}-{m:02d}.xlsx", wb)
        total += len(rows)
        if f"{y}-{m:02d}" >= cur:
            ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
            st["nvd_next"] = cur = f"{ny}-{nm:02d}"
    if total:
        meta_note(sink, f"{TI}/NVD CVE Records", {
            "source": "NIST National Vulnerability Database CVE API 2.0",
            "url": "https://nvd.nist.gov/developers/vulnerabilities",
            "license": "US Government work, public domain",
            "notice": "This product uses the NVD API but is not endorsed or certified by the NVD."})
    return total


def _ext(o):
    for r in o.get("external_references", []):
        if r.get("source_name", "").startswith("mitre") and r.get("external_id"):
            return r["external_id"]
    return ""


def attack(sink, st):
    s = session()
    n = 0
    for dom in ("enterprise-attack", "ics-attack", "mobile-attack"):
        r = get(s, f"https://raw.githubusercontent.com/mitre/cti/master/{dom}/{dom}.json", timeout=300)
        if not r:
            continue
        h = hashlib.md5(r.content).hexdigest()
        if st.get(f"attack_{dom}") == h:
            continue
        objs = [o for o in r.json()["objects"] if not o.get("revoked") and not o.get("x_mitre_deprecated")]
        ids = {o["id"]: o for o in objs}
        tech = [[_ext(o), o["name"], ", ".join(p["phase_name"] for p in o.get("kill_chain_phases", [])),
                 ", ".join(o.get("x_mitre_platforms", [])), o.get("x_mitre_is_subtechnique", False),
                 o.get("x_mitre_detection", "")[:5000], o.get("description", "")[:20000]]
                for o in objs if o["type"] == "attack-pattern"]
        groups = [[_ext(o), o["name"], ", ".join(o.get("aliases", [])), o.get("description", "")[:20000]]
                  for o in objs if o["type"] == "intrusion-set"]
        soft = [[_ext(o), o["name"], o["type"], ", ".join(o.get("x_mitre_platforms", [])),
                 o.get("description", "")[:20000]] for o in objs if o["type"] in ("malware", "tool")]
        mit = [[_ext(o), o["name"], o.get("description", "")[:20000]]
               for o in objs if o["type"] == "course-of-action"]
        rel = []
        for o in objs:
            if o["type"] == "relationship" and o["source_ref"] in ids and o["target_ref"] in ids:
                a, b = ids[o["source_ref"]], ids[o["target_ref"]]
                rel.append([_ext(a), a.get("name"), o["relationship_type"], _ext(b), b.get("name")])
        wb = new_book()
        add_sheet(wb, "Techniques", ["ID", "Name", "Tactics", "Platforms", "Sub-technique", "Detection",
                                     "Description"], sorted(tech), {"Description": 90, "Detection": 60})
        add_sheet(wb, "Groups", ["ID", "Name", "Aliases", "Description"], sorted(groups), {"Description": 90})
        add_sheet(wb, "Software", ["ID", "Name", "Type", "Platforms", "Description"], sorted(soft),
                  {"Description": 90})
        add_sheet(wb, "Mitigations", ["ID", "Name", "Description"], sorted(mit), {"Description": 90})
        add_sheet(wb, "Relationships", ["Source ID", "Source", "Relationship", "Target ID", "Target"], rel)
        base = f"{TI}/MITRE ATT&CK"
        save_book(sink, f"{base}/{dom} {dt.date.today()}.xlsx", wb)
        sink.write_bytes(f"{base}/stix/{dom} {dt.date.today()}.json.gz", gzip.compress(r.content))
        st[f"attack_{dom}"] = h
        n += len(tech)
    if n:
        meta_note(sink, f"{TI}/MITRE ATT&CK", {"source": "MITRE ATT&CK (mitre/cti STIX 2.0)",
                                               "url": "https://github.com/mitre/cti",
                                               "license": "ATT&CK Terms of Use (royalty-free, attribution)"})
    return n


def _prose(part, depth=0):
    out = []
    if part.get("prose"):
        out.append(("  " * depth) + part["prose"])
    for p in part.get("parts", []):
        out += _prose(p, depth + 1)
    return out


def nist(sink, st):
    if st.get("nist"):
        return 0
    s = session()
    url = ("https://raw.githubusercontent.com/usnistgov/oscal-content/main/nist.gov/SP800-53/rev5/json/"
           "NIST_SP-800-53_rev5_catalog.json")
    r = get(s, url, timeout=300)
    if not r:
        return 0
    cat = r.json()["catalog"]
    rows = []
    base = f"{STD}/NIST SP 800-53 Rev 5"
    for g in cat["groups"]:
        doc = Document()
        doc.styles["Normal"].font.size = Pt(10)
        doc.add_heading(f"NIST SP 800-53 Rev. 5 - {g['title']} ({g['id'].upper()})", 0)

        def emit(c, level):
            stmt = [p for p in c.get("parts", []) if p["name"] == "statement"]
            guid = [p for p in c.get("parts", []) if p["name"] == "guidance"]
            text = "\n".join(x for p in stmt for x in _prose(p))
            rows.append([g["title"], c["id"].upper(), c["title"], level, text[:32000],
                         "\n".join(x for p in guid for x in _prose(p))[:32000]])
            doc.add_heading(f"{c['id'].upper()} {c['title']}", 1 if level == "Control" else 2)
            for line in text.split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())
            for p in guid:
                doc.add_paragraph().add_run("Discussion: " + " ".join(_prose(p))).italic = True
            for e in c.get("controls", []):
                emit(e, "Enhancement")

        for c in g.get("controls", []):
            emit(c, "Control")
        p = sink.path(f"{base}/Families/{g['id'].upper()} {g['title']}.docx")
        doc.save(p)
        sink.wrote(p)
    wb = new_book()
    add_sheet(wb, "Controls", ["Family", "ID", "Title", "Level", "Statement", "Discussion"], rows,
              {"Statement": 90, "Discussion": 90, "Title": 40})
    save_book(sink, f"{base}/NIST SP 800-53 Rev 5 control catalog.xlsx", wb)
    meta_note(sink, base, {"source": "NIST SP 800-53 Rev. 5 (OSCAL)", "url": url,
                           "license": "US Government work, public domain"})
    st["nist"] = True
    return len(rows)


FIXED_FILES = [
    ("uci_incident", ITSM + "/ITSM Incident Event Log (UCI)",
     "https://archive.ics.uci.edu/static/public/498/incident+management+process+enriched+event+log.zip",
     "Incident management process enriched event log.zip",
     "UCI Machine Learning Repository: Incident management process enriched event log (ServiceNow)",
     "CC BY 4.0"),
    ("hotels", HOSP + "/Hotel Booking Demand",
     "https://raw.githubusercontent.com/rfordatascience/tidytuesday/master/data/2020/2020-02-11/hotels.csv",
     "hotels.csv", "Hotel booking demand datasets (Antonio, Almeida and Nunes, Data in Brief 2019)",
     "CC BY 4.0"),
]


def fixed(sink, st):
    s = session()
    n = 0
    for key, folder, url, name, title, lic in FIXED_FILES:
        if st.get(key):
            continue
        r = get(s, url, timeout=600)
        if not r:
            continue
        sink.write_bytes(f"{folder}/{name}", r.content)
        meta_note(sink, folder, {"source": title, "url": url, "license": lic})
        st[key] = True
        n += 1
    return n


# ---------- catalogue searches ----------

KEEP_EXT = {"csv", "tsv", "xlsx", "xls", "xlsm", "json", "geojson", "xml", "txt", "log", "pdf", "docx", "doc",
            "pptx", "ppt", "zip", "gz", "parquet", "pcap", "ipynb", "md", "odt", "ods", "rtf", "jsonl"}


def _download(s, url, cap, sink, rel):
    """Stream url to the sink path; drop it if it is bigger than cap. Returns True when saved."""
    p = sink.path(rel)
    try:
        with s.get(url, stream=True, timeout=180) as r:
            if r.status_code != 200 or int(r.headers.get("Content-Length") or 0) > cap:
                return False
            n = 0
            with open(p, "wb") as f:
                for chunk in r.iter_content(1 << 20):
                    n += len(chunk)
                    if n > cap:
                        break
                    f.write(chunk)
        if n > cap or n == 0:
            os.remove(p)
            return False
        sink.wrote(p)
        return True
    except Exception:
        if os.path.exists(p):
            os.remove(p)
        return False


def _open(lic):
    lic = (lic or "").lower()
    return any(lic.startswith(x) or x in lic for x in OPEN_LICENSES) and "closed" not in lic \
        and "nc" not in lic.split("-")


def zenodo(sink, st, per_query=6, cap=2 * 2 ** 30):
    s = session()
    seen = st.setdefault("zenodo_seen", [])
    seen_set = set(seen)
    pages = st.setdefault("zenodo_page", {})
    n = 0
    for folder, qs in SEARCHES.items():
        for q in qs:
            if sink.full():
                return n
            page = pages.get(q, 1)
            time.sleep(2.5)
            r = get(s, "https://zenodo.org/api/records", params={
                "q": q, "size": 25, "page": page, "sort": "mostviewed", "access_status": "open"})
            hits = r.json().get("hits", {}).get("hits", []) if r else []
            pages[q] = page + 1 if hits else 1
            took = 0
            for h in hits:
                if took >= per_query or sink.full():
                    break
                rid = str(h["id"])
                md = h.get("metadata", {})
                lic = (md.get("license") or {}).get("id", "")
                if rid in seen_set or not _open(lic):
                    continue
                seen.append(rid)
                seen_set.add(rid)
                files = [f for f in h.get("files", []) if f["key"].rsplit(".", 1)[-1].lower() in KEEP_EXT
                         and f.get("size", 0) <= cap]
                if not files:
                    continue
                dest = f"{folder}/Zenodo/{safe(md.get('title', rid), 90)} [{rid}]"
                got = 0
                for f in files[:12]:
                    time.sleep(1)
                    if _download(s, f["links"]["self"], cap, sink, f"{dest}/{f['key']}"):
                        got += 1
                if got:
                    meta_note(sink, dest, {"source": "Zenodo", "record": h.get("links", {}).get("self_html"),
                                           "doi": h.get("doi"), "title": md.get("title"),
                                           "license": lic, "search": q,
                                           "description": (md.get("description") or "")[:4000]})
                    took += 1
                    n += 1
    return n


def datagov(sink, st, per_query=5, cap=1 * 2 ** 30):
    s = session()
    seen = st.setdefault("datagov_seen", [])
    seen_set = set(seen)
    offs = st.setdefault("datagov_start", {})
    n = 0
    for folder, qs in SEARCHES.items():
        for q in qs:
            if sink.full():
                return n
            start = offs.get(q, 0)
            time.sleep(1)
            r = get(s, "https://catalog.data.gov/api/3/action/package_search",
                    params={"q": q, "rows": 20, "start": start})
            res = r.json().get("result", {}).get("results", []) if r else []
            offs[q] = start + 20 if res else 0
            took = 0
            for p in res:
                if took >= per_query or sink.full():
                    break
                if p["id"] in seen_set or not _open(p.get("license_id") or "public"):
                    continue
                seen.append(p["id"])
                seen_set.add(p["id"])
                res_ok = [x for x in p.get("resources", [])
                          if (x.get("format") or "").lower().strip(".") in KEEP_EXT and x.get("url")]
                if not res_ok:
                    continue
                dest = f"{folder}/data.gov/{safe(p.get('title', p['name']), 90)}"
                got = 0
                for x in res_ok[:6]:
                    ext = (x.get("format") or "dat").lower().strip(".")
                    nm = safe(x.get("name") or p["name"], 80)
                    if not nm.lower().endswith("." + ext):
                        nm += "." + ext
                    if _download(s, x["url"], cap, sink, f"{dest}/{nm}"):
                        got += 1
                if got:
                    meta_note(sink, dest, {"source": "data.gov", "dataset": f"https://catalog.data.gov/dataset/{p['name']}",
                                           "publisher": (p.get("organization") or {}).get("title"),
                                           "license": p.get("license_title"), "search": q,
                                           "description": (p.get("notes") or "")[:4000]})
                    took += 1
                    n += 1
    return n


SOURCES = {"kev": kev, "nvd": nvd, "attack": attack, "nist": nist, "fixed": fixed,
           "zenodo": zenodo, "datagov": datagov}
