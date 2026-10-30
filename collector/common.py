"""Shared helpers: HTTP session, run state, staging sink that ships files to Drive with rclone."""
import json
import os
import re
import shutil
import subprocess
import time

import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

UA = "ops-data-collector/1.0 (research; github.com/Elsavado/ops-data-collector)"


def session():
    s = requests.Session()
    s.headers["User-Agent"] = UA
    return s


def get(s, url, tries=5, **kw):
    """GET with backoff on 429/5xx and network errors. Returns the response or None."""
    kw.setdefault("timeout", 120)
    for i in range(tries):
        try:
            r = s.get(url, **kw)
            if r.status_code == 200:
                return r
            if r.status_code in (403, 404, 410):
                return None
        except requests.RequestException:
            pass
        time.sleep(min(10 * 2 ** i, 120))
    return None


def safe(name, limit=120):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", str(name))
    name = re.sub(r"\s+", " ", name).strip(" .")
    return (name[:limit].rstrip(" .") or "untitled")


class State:
    def __init__(self, path):
        self.path = path
        self.data = {}
        if os.path.exists(path):
            with open(path, encoding="utf-8") as f:
                self.data = json.load(f)

    def section(self, key):
        return self.data.setdefault(key, {})

    def save(self):
        os.makedirs(os.path.dirname(self.path) or ".", exist_ok=True)
        tmp = self.path + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.data, f, indent=1, sort_keys=True)
        os.replace(tmp, self.path)


class Sink:
    """Files are written under out/<drive path>. flush() moves them to the Drive root with rclone
    (live) or just counts and deletes them (dry run). Logs only show counts, never names."""

    def __init__(self, out, remote, live, state, budget_bytes, deadline, keep=False):
        self.out, self.remote, self.live, self.state, self.keep = out, remote, live, state, keep
        self.budget = budget_bytes
        self.deadline = deadline
        self.slice_end = deadline  # per-source time slice, set by the runner
        self.sent_bytes = 0
        self.files = 0
        self.pending = 0
        os.makedirs(out, exist_ok=True)

    def path(self, rel):
        p = os.path.join(self.out, *[safe(x, 150) for x in rel.split("/")])
        os.makedirs(os.path.dirname(p), exist_ok=True)
        return p

    def wrote(self, p):
        n = os.path.getsize(p)
        self.pending += n
        self.files += 1
        if self.pending > 400 * 2 ** 20:
            self.flush()

    def write_bytes(self, rel, data):
        p = self.path(rel)
        with open(p, "wb") as f:
            f.write(data)
        self.wrote(p)

    def write_text(self, rel, text):
        self.write_bytes(rel, text.encode("utf-8"))

    def full(self):
        return self.sent_bytes + self.pending >= self.budget or time.time() > min(self.deadline, self.slice_end)

    def flush(self):
        if self.keep and not self.live:
            self.state.save()
            self.sent_bytes += self.pending
            self.pending = 0
            return
        if self.pending == 0 and not os.listdir(self.out):
            self.state.save()
            return
        if self.live:
            cmd = ["rclone", "move", self.out, self.remote, "--transfers", "4", "--checkers", "8",
                   "--drive-chunk-size", "64M", "--retries", "5", "--low-level-retries", "20",
                   "--delete-empty-src-dirs", "-q"]
            subprocess.run(cmd, check=True)
            self.state.save()
            subprocess.run(["rclone", "copyto", self.state.path,
                            self.remote + "/_state/" + os.path.basename(self.state.path), "-q"],
                           check=True)
        else:
            self.state.save()
            shutil.rmtree(self.out, ignore_errors=True)
        os.makedirs(self.out, exist_ok=True)
        self.sent_bytes += self.pending
        self.pending = 0


HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def add_sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(safe(title, 31)[:31])
    ws.append(list(header))
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        w = (widths or {}).get(h) or min(max(len(str(h)) + 2, 12), 50)
        ws.column_dimensions[get_column_letter(i)].width = w
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def new_book():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def save_book(sink, rel, wb, title=None, note=None):
    if title or note:
        wb.properties.title = title
        wb.properties.description = note
    p = sink.path(rel)
    wb.save(p)
    sink.wrote(p)
