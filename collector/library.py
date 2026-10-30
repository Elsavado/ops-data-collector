"""Books & knowledge library: open-access books, NIST guides, Wikipedia topic articles, OWASP, CWE/CAPEC."""
import datetime as dt
import html
import re
import time

from docx import Document
from docx.shared import Pt

from .common import get, safe, session
from .public import LIB, _download, meta_note

TOPICS = ["cybersecurity", "information security", "security operations center", "incident response",
          "video surveillance", "CCTV", "physical security", "access control", "IT service management",
          "ITIL", "cloud computing", "network security", "risk management", "business continuity",
          "facility management", "hospitality management", "hotel management", "cruise tourism",
          "maritime security", "port security", "security management", "digital forensics",
          "data protection privacy", "internet of things security", "critical infrastructure protection",
          "managed services", "service level agreement", "security guard", "crime prevention"]

BOOK_LICENSES = ("cc-by", "cc-by-sa", "cc-by-nd", "cc-by-nc", "cc-by-nc-sa", "cc-by-nc-nd", "cc0", "public-domain")

NIST_SP = """800-12r1 800-30r1 800-34r1 800-37r2 800-39 800-40r4 800-41r1 800-46r2 800-50r1 800-53r5 800-53Ar5
800-53B 800-55v1 800-60v1r1 800-61r2 800-61r3 800-63-3 800-63B 800-82r3 800-83r1 800-84 800-86 800-88r1 800-92
800-94 800-100 800-115 800-116r1 800-123 800-124r2 800-125 800-128 800-137 800-144 800-145 800-146 800-150
800-153 800-160v1r1 800-160v2r1 800-161r1 800-171r3 800-172 800-181r1 800-184 800-190 800-207 800-209 800-210
800-213 800-215 800-218 800-221 800-222""".split()
NIST_OTHER = ["https://nvlpubs.nist.gov/nistpubs/CSWP/NIST.CSWP.29.pdf",
              "https://nvlpubs.nist.gov/nistpubs/ir/2018/NIST.IR.8183.pdf",
              "https://nvlpubs.nist.gov/nistpubs/ir/2021/NIST.IR.8286.pdf",
              "https://nvlpubs.nist.gov/nistpubs/ir/2020/NIST.IR.8259.pdf",
              "https://nvlpubs.nist.gov/nistpubs/ir/2016/NIST.IR.8112.pdf"]

WIKI_CATS = {
    "Security Operations": ["Computer security", "Computer network security", "Intrusion detection systems",
                            "Security information and event management", "Computer security procedures",
                            "Incident management", "Malware", "Cyberattacks", "Computer forensics"],
    "Surveillance & Physical Security": ["Video surveillance", "Closed-circuit television", "Physical security",
                                         "Access control", "Security engineering", "Alarms", "Locksmithing",
                                         "Biometrics", "Security guards"],
    "IT Service Management": ["IT service management", "Information technology management", "Help desk",
                              "Configuration management", "Network management", "System administration"],
    "Cloud & Infrastructure": ["Cloud computing", "Data centers", "Computer networking", "Video storage",
                               "Network video recorders", "Power over Ethernet", "Internet of things"],
    "Hospitality & Maritime": ["Hospitality management", "Hotel terminology", "Cruise ships", "Maritime security",
                               "Ports and harbours", "Tourism"],
    "IT & Developer Tools": ["Software development tools", "Version control systems", "Continuous integration",
                             "Configuration management", "Integrated development environments", "Linux distributions",
                             "Operating systems", "Database management systems", "Virtualization software",
                             "Containerization software", "Network protocols", "Web servers", "Programming languages",
                             "Monitoring software", "Remote administration software", "Backup software",
                             "Identity management", "Firewall software", "Software testing", "DevOps"],
    "Business & Risk": ["Risk management", "Business continuity", "Service-level agreements",
                        "Outsourcing", "Facility management", "Security companies", "Privacy law"],
}


def nist_pubs(sink, st):
    s = session()
    done = set(st.setdefault("nist_pubs", []))
    urls = [f"https://nvlpubs.nist.gov/nistpubs/SpecialPublications/NIST.SP.{x}.pdf" for x in NIST_SP] + NIST_OTHER
    n = 0
    for u in urls:
        if u in done or sink.full():
            continue
        name = u.rsplit("/", 1)[1]
        time.sleep(1)
        if _download(s, u, 200 * 2 ** 20, sink, f"{LIB}/NIST Security Publications/{name}"):
            n += 1
        st["nist_pubs"].append(u)
    if n:
        meta_note(sink, f"{LIB}/NIST Security Publications", {
            "source": "NIST Special Publications and reports", "url": "https://csrc.nist.gov/publications",
            "license": "US Government work, public domain"})
    return n


def open_books(sink, st, per_topic=8):
    """Open-access books (CC or public domain, with a PDF) found through OpenAlex."""
    s = session()
    seen = st.setdefault("books_seen", [])
    seen_set = set(seen)
    pages = st.setdefault("books_page", {})
    n = 0
    for topic in TOPICS:
        if sink.full():
            break
        page = pages.get(topic, 1)
        time.sleep(1)
        r = get(s, "https://api.openalex.org/works", params={
            "filter": f"title.search:{topic},type:book|book-chapter,open_access.is_oa:true",
            "per-page": 50, "page": page, "sort": "cited_by_count:desc",
            "select": "id,doi,title,publication_year,authorships,best_oa_location,primary_location,type"})
        res = r.json().get("results", []) if r else []
        pages[topic] = page + 1 if res else 1
        took = 0
        for w in res:
            if took >= per_topic or sink.full():
                break
            loc = w.get("best_oa_location") or {}
            lic, pdf = (loc.get("license") or ""), loc.get("pdf_url")
            if w["id"] in seen_set or not pdf or lic not in BOOK_LICENSES:
                continue
            seen.append(w["id"])
            seen_set.add(w["id"])
            title = safe(w.get("title") or w["id"], 110)
            kind = "Books" if w.get("type") == "book" else "Book chapters"
            dest = f"{LIB}/Open Access Books/{safe(topic.title(), 40)}/{kind}"
            time.sleep(1)
            fname = f"{title} ({w.get('publication_year')})"
            if _download(s, pdf, 300 * 2 ** 20, sink, f"{dest}/{fname}.pdf"):
                meta_note(sink, f"{dest}/{fname}", {
                    "title": w.get("title"), "doi": w.get("doi"), "year": w.get("publication_year"),
                    "authors": [a["author"]["display_name"] for a in w.get("authorships", [])][:20],
                    "publisher": ((loc.get("source") or {}).get("display_name")), "license": lic,
                    "pdf": pdf, "found_via": "OpenAlex"})
                took += 1
                n += 1
    return n


def _wiki_doc(title, html_text, url, credit="Wikipedia contributors, CC BY-SA 4.0"):
    d = Document()
    d.styles["Normal"].font.size = Pt(10.5)
    d.core_properties.title = title
    d.core_properties.comments = f"{url} - {credit}"
    d.add_heading(title, 0)
    d.add_paragraph(f"Source: {url} ({credit}). Retrieved {dt.date.today()}.")
    body = re.sub(r"(?s)<(style|script|table|figure|sup)[^>]*>.*?</\1>", "", html_text)
    for tag, inner in re.findall(r"(?s)<(h2|h3|h4|p|li)[^>]*>(.*?)</\1>", body):
        text = html.unescape(re.sub(r"<[^>]+>", "", inner)).strip()
        text = re.sub(r"\[\d+\]", "", text)
        if not text:
            continue
        if tag == "h2":
            d.add_heading(text, 1)
        elif tag in ("h3", "h4"):
            d.add_heading(text, 2)
        elif tag == "li":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    return d


def wikipedia(sink, st, per_cat=40):
    s = session()
    seen = set(st.setdefault("wiki_seen", []))
    cont = st.setdefault("wiki_cont", {})
    api = "https://en.wikipedia.org/w/api.php"
    n = 0
    for group, cats in WIKI_CATS.items():
        for cat in cats:
            if sink.full() or cont.get(cat) == "done":
                continue
            params = {"action": "query", "list": "categorymembers", "cmtitle": f"Category:{cat}",
                      "cmtype": "page", "cmlimit": per_cat, "format": "json"}
            if cont.get(cat):
                params["cmcontinue"] = cont[cat]
            r = get(s, api, params=params)
            if not r:
                continue
            d = r.json()
            cont[cat] = d.get("continue", {}).get("cmcontinue", "done")
            for m in d.get("query", {}).get("categorymembers", []):
                t = m["title"]
                if t in seen or sink.full():
                    continue
                time.sleep(0.2)
                pr = get(s, f"https://en.wikipedia.org/api/rest_v1/page/html/{t.replace(' ', '_')}")
                if not pr:
                    continue
                url = f"https://en.wikipedia.org/wiki/{t.replace(' ', '_')}"
                doc = _wiki_doc(t, pr.text, url)
                p = sink.path(f"{LIB}/Wikipedia Knowledge Base/{group}/{safe(cat, 50)}/{safe(t, 100)}.docx")
                doc.save(p)
                sink.wrote(p)
                seen.add(t)
                st["wiki_seen"].append(t)
                n += 1
    if n:
        meta_note(sink, f"{LIB}/Wikipedia Knowledge Base", {
            "source": "English Wikipedia", "license": "CC BY-SA 4.0 (Wikipedia contributors)",
            "note": "Articles converted to DOCX; text only, tables and images left out."})
    return n


KNOWLEDGE = [
    ("owasp", "OWASP Cheat Sheet Series", "https://github.com/OWASP/CheatSheetSeries/archive/refs/heads/master.zip",
     "OWASP Cheat Sheet Series.zip", "CC BY-SA 4.0"),
    ("owasp_asvs", "OWASP ASVS", "https://github.com/OWASP/ASVS/archive/refs/heads/master.zip",
     "OWASP Application Security Verification Standard.zip", "CC BY-SA 4.0"),
    ("cwe", "MITRE CWE", "https://cwe.mitre.org/data/xml/cwec_latest.xml.zip", "CWE catalog (latest).xml.zip",
     "CWE Terms of Use (free, attribution)"),
    ("capec", "MITRE CAPEC", "https://capec.mitre.org/data/xml/capec_latest.xml", "CAPEC catalog (latest).xml",
     "CAPEC Terms of Use (free, attribution)"),
    ("d3fend", "MITRE D3FEND", "https://d3fend.mitre.org/ontologies/d3fend.json", "D3FEND ontology.json",
     "MIT licence"),
    ("sigma", "SigmaHQ detection rules", "https://github.com/SigmaHQ/sigma/archive/refs/heads/master.zip",
     "Sigma detection rules.zip", "Detection Rule License 1.1"),
    ("atomic", "Atomic Red Team tests", "https://github.com/redcanaryco/atomic-red-team/archive/refs/heads/master.zip",
     "Atomic Red Team.zip", "MIT licence"),
]


def knowledge(sink, st):
    s = session()
    n = 0
    for key, title, url, name, lic in KNOWLEDGE:
        if sink.full():
            break
        stamp = st.get(f"kb_{key}")
        if stamp and stamp > (dt.date.today() - dt.timedelta(days=90)).isoformat():
            continue  # refresh each quarter
        if _download(s, url, 2 * 2 ** 30, sink, f"{LIB}/Security Knowledge Bases/{title}/{dt.date.today()} {name}"):
            meta_note(sink, f"{LIB}/Security Knowledge Bases/{title}", {"source": title, "url": url, "license": lic})
            st[f"kb_{key}"] = dt.date.today().isoformat()
            n += 1
    return n


SOURCES = {"nistpubs": nist_pubs, "books": open_books, "wikipedia": wikipedia, "knowledge": knowledge}
