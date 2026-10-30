"""Fictional security & IT managed-services firm "Harborlight Secure Systems".

Everything here is SYNTHETIC: the company, its clients, people, suppliers and numbers are made up.
The world (clients, sites, devices, staff) is simulated month by month from a fixed seed, so a rerun
reproduces past months exactly and only appends new ones. Month files carry their own seeded RNG.
"""
import datetime as dt
import gzip
import io
import math
import random
import sys
import zipfile
from email.message import EmailMessage
from email.utils import format_datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import names as N
from . import staff_records as S
from .common import add_sheet, new_book, safe, save_book

CO = "Harborlight Secure Systems"
DOMAIN = "harborlight.example"
ROOT = f"Fictional Company - {CO}"
F = {k: f"{ROOT}/{v}" for k, v in {
    "crm": "01 Clients & CRM", "contracts": "02 Contracts", "soc": "03 SOC Monitoring",
    "inc": "04 Incidents", "field": "05 Field Service & Maintenance", "help": "06 IT Help Desk",
    "assets": "07 Assets & Inventory", "fin": "08 Finance & Accounting", "hr": "09 HR & Workforce",
    "rep": "10 Management Reports", "pol": "11 Policies & SOPs", "mail": "12 Emails & Meetings"}.items()}
BANNER = "FICTIONAL - synthetic research data. Harborlight Secure Systems and all names are made up."
START = (2016, 1)
SEED = 20160101

ACCOUNTS = {
    "1000": ("Cash at bank", "Asset"), "1100": ("Accounts receivable", "Asset"),
    "1500": ("Equipment & vehicles", "Asset"), "1510": ("Accumulated depreciation", "Asset"),
    "2000": ("Accounts payable", "Liability"), "2100": ("VAT / sales tax payable", "Liability"),
    "2200": ("End-of-service benefit provision", "Liability"), "3000": ("Share capital", "Equity"),
    "3100": ("Retained earnings", "Equity"), "4000": ("Managed monitoring revenue", "Revenue"),
    "4100": ("Installation revenue", "Revenue"), "4200": ("Parts & materials revenue", "Revenue"),
    "5000": ("Salaries & allowances", "Expense"), "5050": ("End-of-service benefit expense", "Expense"),
    "5100": ("Parts & materials cost", "Expense"), "5200": ("Cloud hosting", "Expense"),
    "5300": ("Telecom & connectivity", "Expense"), "5400": ("Rent", "Expense"),
    "5500": ("Vehicles", "Expense"), "5600": ("Software licences", "Expense"),
    "5700": ("Insurance", "Expense"), "5800": ("Depreciation", "Expense"),
    "5900": ("Bad debts", "Expense"), "6000": ("Training, uniforms & PPE", "Expense"),
    "6100": ("Professional fees", "Expense"), "6200": ("Staff travel & expense claims", "Expense")}

# role: (department, monthly base USD, required-headcount rule on (cameras, sites, clients))
ROLES = {
    "Chief Executive Officer": ("Executive", 16000, lambda c, s, k: 1),
    "Operations Director": ("Executive", 12500, lambda c, s, k: 1),
    "Finance Manager": ("Finance", 8200, lambda c, s, k: 1),
    "SOC Manager": ("Security Operations", 8800, lambda c, s, k: 1),
    "Field Services Manager": ("Field Services", 7600, lambda c, s, k: 1),
    "SOC Shift Lead": ("Security Operations", 4300, lambda c, s, k: 4 + c // 6000 * 4),
    "SOC Analyst": ("Security Operations", 2700, lambda c, s, k: 8 + c // 450),
    "Field Technician": ("Field Services", 2400, lambda c, s, k: 2 + c // 500),
    "Network Engineer": ("Technology", 5000, lambda c, s, k: 1 + s // 14),
    "Cloud Engineer": ("Technology", 5900, lambda c, s, k: 1 + k // 18),
    "IT Support Specialist": ("Technology", 2900, lambda c, s, k: 1 + k // 9),
    "Account Manager": ("Commercial", 5400, lambda c, s, k: 1 + k // 12),
    "Project Manager": ("Projects", 6000, lambda c, s, k: 1 + k // 20),
    "Finance Officer": ("Finance", 3600, lambda c, s, k: 1 + k // 25),
    "HR Officer": ("People", 3400, lambda c, s, k: 1 + k // 35),
}
HEADS = {"Executive": "Chief Executive Officer", "Finance": "Finance Manager",
         "Security Operations": "SOC Manager", "Field Services": "Field Services Manager",
         "Technology": "Operations Director", "Commercial": "Chief Executive Officer",
         "Projects": "Operations Director", "People": "Chief Executive Officer"}

DEVICES = [  # kind, model, unit cost, share per camera
    ("Camera", "4MP IR dome IP camera", 210, 0.55), ("Camera", "4K bullet IP camera", 340, 0.25),
    ("Camera", "360 fisheye IP camera", 520, 0.12), ("Camera", "PTZ 30x IP camera", 1450, 0.08),
    ("NVR", "64-channel network video recorder", 3900, 1 / 48), ("Access door", "Card reader + controller", 640, 0.12),
    ("Switch", "24-port PoE+ managed switch", 1100, 1 / 22), ("Server", "Video analytics server", 7800, 1 / 400),
]

ALERTS = [  # type, weight, severity, escalation chance
    ("Motion in restricted zone after hours", 18, "Medium", .010), ("Camera offline", 14, "Low", 0),
    ("Video loss", 6, "Low", 0), ("Camera tampering / view obstructed", 5, "Medium", .02),
    ("Intrusion - perimeter line crossed", 7, "High", .04), ("Loitering detected", 9, "Low", .005),
    ("Door forced open", 4, "High", .05), ("Door held open", 10, "Low", .002),
    ("Invalid access card attempts", 6, "Medium", .01), ("Smoke / fire analytic", 1.2, "Critical", .2),
    ("Crowd density threshold", 3, "Medium", .01), ("Abandoned object", 2, "High", .06),
    ("NVR disk failure", 1.5, "Medium", 0), ("Network switch down", 1.5, "High", 0),
    ("Failed admin logins on VMS", 2, "High", .05), ("Malware signature on recorder", .3, "Critical", .5),
    ("Unauthorised deck / crew area access", 2.5, "High", .04), ("Panic button activated", .6, "Critical", .3),
]
DISPOSITIONS = ["False positive - environmental", "False positive - authorised activity",
                "Verified - no action required", "Verified - client security notified",
                "Verified - guard dispatched"]
TICKET_CATS = [("VMS login / password", "P3"), ("Footage export request", "P2"),
               ("Camera view change request", "P4"), ("New user account", "P4"),
               ("Network connectivity issue", "P2"), ("Access card issue", "P3"),
               ("Report request", "P4"), ("Cloud retention extension", "P3"),
               ("Mobile app access", "P3"), ("Live view slow / buffering", "P3"),
               ("Recording gap investigation", "P2"), ("Firmware update request", "P4")]
SLA = {"P1": (15, 4), "P2": (30, 8), "P3": (120, 24), "P4": (480, 72)}  # first response min, resolve h
SHIFTS = [("Day", 6), ("Evening", 14), ("Night", 22)]


def ym_add(ym, n):
    y, m = ym
    t = y * 12 + m - 1 + n
    return (t // 12, t % 12 + 1)


def ym_s(ym):
    return f"{ym[0]}-{ym[1]:02d}"


def mdays(ym):
    nxt = ym_add(ym, 1)
    return (dt.date(nxt[0], nxt[1], 1) - dt.date(ym[0], ym[1], 1)).days


def pois(rng, mu):
    if mu <= 0:
        return 0
    if mu > 30:
        return max(0, int(round(rng.gauss(mu, math.sqrt(mu)))))
    L, k, p = math.exp(-mu), 0, 1.0
    while True:
        p *= rng.random()
        if p < L:
            return k
        k += 1


def r2(x):
    return round(x + 1e-9, 2)


# ---------------------------------------------------------------- world simulation

class World:
    def __init__(self, end, scale=1.0, seed=SEED):
        rng = random.Random(seed)
        self.scale, self.end = scale, end
        self.clients, self.sites, self.contracts, self.staff, self.assets = [], [], [], [], []
        self.failures = {}
        self.prospects_lost = {}
        used_names, used_people = set(), set()
        self._rng = rng

        def person():
            while True:
                n = f"{rng.choice(N.FIRST)} {rng.choice(N.LAST)}"
                if n not in used_people:
                    used_people.add(n)
                    return n

        self.person = person
        ym, i = START, 0
        while ym <= end:
            # new clients (demand grows over time)
            mu = (4 if i == 0 else 0.35 + 0.006 * i) * scale
            for _ in range(pois(rng, mu)):
                self._new_client(rng, ym, used_names, person)
            # contract renewals / churn
            for c in list(self.contracts):
                if c["end"] == ym and c["status"] == "Active":
                    cl = self.clients[c["client"]]
                    if rng.random() < 0.82:
                        c["status"] = "Renewed"
                        self._contract(rng, cl, ym_add(ym, 1), c["fee_per_camera"] * rng.uniform(1.02, 1.06),
                                       "Renewal")
                    else:
                        c["status"] = "Expired"
                        cl["end"] = ym_add(ym, 1)
            # device failures and replacements
            fails = []
            for a in self.assets:
                if a["status"] == "In service" and a["installed"] < ym:
                    age = (ym[0] - a["installed"][0]) * 12 + ym[1] - a["installed"][1]
                    if self.clients[self.sites[a["site"]]["client"]]["end"] and \
                            self.clients[self.sites[a["site"]]["client"]]["end"] <= ym:
                        a["status"], a["retired"] = "Removed (contract ended)", ym
                        continue
                    if rng.random() < 0.0025 + 0.00012 * age:
                        a["status"], a["retired"] = "Failed - replaced", ym
                        new = self._asset(rng, a["site"], a["kind"], a["model"], a["cost"], ym)
                        fails.append((a["id"], new["id"]))
            self.failures[ym] = fails
            # staffing
            cams, sites, clients = self.load(ym)
            for e in self.staff:
                if e["left"] is None and e["role"] not in HEADS.values() and rng.random() < 0.011:
                    e["left"] = ym
                    e["leave_reason"] = rng.choice(["Resigned", "Resigned", "Relocated", "Contract ended",
                                                    "Dismissed"])
            for role, (dept, base, rule) in ROLES.items():
                need = rule(cams, sites, clients)
                have = sum(1 for e in self.staff if e["role"] == role and e["left"] is None)
                for _ in range(max(0, need - have)):
                    self._hire(rng, role, dept, base, ym, person)
            ym, i = ym_add(ym, 1), i + 1
        for e in self.staff:
            e["manager"] = self._manager(e)

    # -- builders
    def _new_client(self, rng, ym, used, person):
        kind = rng.choices(["Hotel", "Cruise line", "Marina / port", "Residential & retail"], [55, 18, 12, 15])[0]
        while True:
            a = rng.choice(N.HOTEL_A)
            b = rng.choice({"Hotel": N.HOTEL_B, "Cruise line": N.CRUISE_B, "Marina / port": N.PORT_B,
                            "Residential & retail": N.TOWER_B}[kind])
            name = f"{a} {b}"
            if name not in used:
                used.add(name)
                break
        city, country, vat = N.location(rng)
        cid = len(self.clients)
        cl = {"id": cid, "code": f"CL{cid + 1:04d}", "name": name, "kind": kind, "city": city,
              "country": country, "vat": vat, "since": ym, "end": None,
              "domain": safe(a.lower(), 30) + ".example", "contact": person(),
              "terms": rng.choice([30, 30, 45, 60]), "sites": []}
        self.clients.append(cl)
        nsites = {"Hotel": rng.choice([1, 1, 1, 2, 3]), "Cruise line": rng.randint(1, 5),
                  "Marina / port": rng.choice([1, 2]), "Residential & retail": rng.choice([1, 2])}[kind]
        for j in range(nsites):
            if kind == "Cruise line":
                sname, cams = f"MS {a} {rng.choice(N.SHIP)}", rng.randint(160, 620)
            else:
                c2 = N.location(rng) if j else (city, country, vat)
                sname = f"{name} - {c2[0]}" if nsites > 1 or kind != "Hotel" else name
                cams = {"Hotel": rng.randint(70, 420), "Marina / port": rng.randint(60, 260),
                        "Residential & retail": rng.randint(90, 380)}[kind]
            sid = len(self.sites)
            st = {"id": sid, "code": f"ST{sid + 1:04d}", "client": cid, "name": sname, "cameras": cams,
                  "live": ym_add(ym, 1)}
            self.sites.append(st)
            cl["sites"].append(sid)
            for kind_, model, cost, share in DEVICES:
                n = max(1, round(cams * share)) if kind_ != "Camera" else 0
                if kind_ == "Camera":
                    n = round(cams * share)
                for _ in range(n):
                    self._asset(rng, sid, kind_, model, cost, ym_add(ym, 1))
        self._contract(rng, cl, ym_add(ym, 1), rng.uniform(26, 44), "New")

    def _contract(self, rng, cl, start, rate, kind):
        term = rng.choice([12, 24, 36, 36])
        cams = sum(self.sites[s]["cameras"] for s in cl["sites"])
        c = {"id": len(self.contracts), "code": f"MSA-{start[0]}-{len(self.contracts) + 1:04d}",
             "client": cl["id"], "start": start, "end": ym_add(start, term - 1), "term": term,
             "fee_per_camera": r2(rate), "base_fee": r2(rng.choice([1500, 2500, 4000]) * (1 + cams / 1500)),
             "status": "Active", "type": kind, "cameras": cams,
             "install_fee": r2(cams * rng.uniform(310, 420)) if kind == "New" else 0.0}
        c["monthly"] = r2(c["base_fee"] + cams * c["fee_per_camera"])
        self.contracts.append(c)

    def _asset(self, rng, sid, kind, model, cost, ym):
        a = {"id": f"AS{len(self.assets) + 1:07d}", "site": sid, "kind": kind, "model": model,
             "cost": cost, "installed": ym, "status": "In service", "retired": None,
             "serial": f"{kind[:2].upper()}{rng.randrange(16 ** 8):08X}",
             "ip": f"10.{sid % 250}.{rng.randint(1, 254)}.{rng.randint(2, 254)}" if kind != "Access door" else ""}
        self.assets.append(a)
        return a

    def _hire(self, rng, role, dept, base, ym, person):
        n = person()
        first, last = n.split(" ", 1)
        e = {"id": f"E{len(self.staff) + 1:05d}", "name": n,
             "email": f"{first.lower()}.{last.lower().replace(' ', '').replace('-', '')}@{DOMAIN}",
             "role": role, "dept": dept, "hired": ym, "left": None, "leave_reason": "",
             "base": r2(base * rng.uniform(0.9, 1.12) * (1.03 ** ((ym[0] - 2016)))),
             "nationality": rng.choice(N.NATIONALITIES),
             "crew": len([x for x in self.staff if x["role"] == role]) % 4}
        self.staff.append(e)

    def _manager(self, e):
        head = HEADS.get(e["dept"])
        if e["role"] == "Chief Executive Officer":
            return ""
        if e["role"] == head:
            return self.staff[0]["id"]
        return next((x["id"] for x in self.staff if x["role"] == head), "")

    # -- queries
    def active_clients(self, ym):
        return [c for c in self.clients if c["since"] < ym and (c["end"] is None or c["end"] > ym)]

    def active_sites(self, ym):
        live = {c["id"] for c in self.active_clients(ym)}
        return [s for s in self.sites if s["client"] in live and s["live"] <= ym]

    def load(self, ym):
        s = self.active_sites(ym)
        return sum(x["cameras"] for x in s), len(s), len(self.active_clients(ym))

    def contract_for(self, cid, ym):
        return next((c for c in self.contracts if c["client"] == cid and c["start"] <= ym <= c["end"]), None)

    def staff_at(self, ym, role=None):
        return [e for e in self.staff if e["hired"] <= ym and (e["left"] is None or e["left"] > ym)
                and (role is None or e["role"] == role)]

    def assets_at(self, ym, sid=None):
        return [a for a in self.assets if a["installed"] <= ym and (a["retired"] is None or a["retired"] > ym)
                and (sid is None or a["site"] == sid)]

    def pay(self, e, ym):
        years = ym[0] - e["hired"][0]
        return r2(e["base"] * 1.03 ** years)


# ---------------------------------------------------------------- docx / pdf helpers

def new_doc(title, subtitle=None):
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    d.core_properties.author = CO
    d.core_properties.comments = BANNER
    sec = d.sections[0]
    p = sec.header.paragraphs[0]
    p.text = f"{CO} | Confidential (fictional)"
    p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(8), RGBColor(0x55, 0x55, 0x55)
    f = sec.footer.paragraphs[0]
    f.text = BANNER
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.runs[0].font.size = Pt(7)
    d.add_heading(title, 0)
    if subtitle:
        d.add_paragraph(subtitle).runs[0].italic = True
    return d


def doc_table(d, header, rows):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = str(h)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v))
    return t


def save_doc(sink, rel, d):
    p = sink.path(rel)
    d.save(p)
    sink.wrote(p)


def bullet(d, text):
    d.add_paragraph(text, style="List Bullet")


# ---------------------------------------------------------------- month generator

class Month:
    def __init__(self, w, ym, ledger, sink, detail=1.0):
        self.w, self.ym, self.L, self.sink, self.detail = w, ym, ledger, sink, detail
        self.rng = random.Random(f"{SEED}-{ym_s(ym)}")
        self.tag = ym_s(ym)
        self.y = str(ym[0])
        self.days = mdays(ym)
        self.d0 = dt.date(ym[0], ym[1], 1)
        self.clients = w.active_clients(ym)
        self.sites = w.active_sites(ym)
        self.site_by = {s["id"]: s for s in self.sites}
        self.cl_by = {c["id"]: c for c in w.clients}
        self.staff = w.staff_at(ym)
        self.by_site = {}
        for a in w.assets_at(ym):
            self.by_site.setdefault(a["site"], []).append(a)
        self.gl = []
        self.emails = []
        self.C, self.F, self.BANNER = sys.modules[__name__], F, BANNER

    def ts(self, day=None, hour=None):
        r = self.rng
        d = self.d0 + dt.timedelta(days=day if day is not None else r.randrange(self.days))
        return dt.datetime.combine(d, dt.time(hour if hour is not None else r.randrange(24),
                                              r.randrange(60), r.randrange(60)))

    def people(self, role):
        return [e for e in self.staff if e["role"] == role] or [e for e in self.staff][:1]

    def on_shift(self, t):
        h = t.hour
        s = 0 if 6 <= h < 14 else 1 if 14 <= h < 22 else 2
        day = (t.date() - dt.date(2016, 1, 1)).days - (1 if s == 2 and h < 6 else 0)
        crew = (day + s) % 4
        an = [e for e in self.people("SOC Analyst") if e["crew"] == crew] or self.people("SOC Analyst")
        return self.rng.choice(an)

    def post(self, date, ref, desc, lines):
        assert abs(sum(d - c for _, d, c in lines)) < 0.01, (ref, lines)
        for acct, dr, cr in lines:
            if dr or cr:
                self.gl.append([date, ref, acct, ACCOUNTS[acct][0], desc, r2(dr), r2(cr)])

    def mail(self, t, frm, to, subject, body):
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = frm, to, subject
        m["Date"] = format_datetime(t.replace(tzinfo=dt.timezone.utc))
        m["Message-ID"] = f"<{self.tag}.{len(self.emails) + 1}@{DOMAIN}>"
        m["X-Dataset-Notice"] = BANNER
        m.set_content(body)
        self.emails.append((t, subject, m.as_bytes()))

    # -------- generators
    def run(self):
        self.soc()
        self.incidents()
        self.field()
        self.helpdesk()
        self.crm()
        self.hr()
        self.finance()
        S.collect(self)
        if self.ym[1] == 12:
            S.person_year(self.w, self.L, self.ym[0], self.sink, self.C, final=True)
        self.reports()
        self.meetings()
        S.meetings_av(self)
        self.write_mail()
        if self.detail > 0:
            self.telemetry()

    def soc(self):
        r, rows, health = self.rng, [], []
        types = [a[0] for a in ALERTS]
        wts = [a[1] for a in ALERTS]
        meta = {a[0]: a for a in ALERTS}
        self.alerts = []
        for s in self.sites:
            ship = s["name"].startswith("MS ")
            devs = [a for a in self.by_site.get(s["id"], []) if a["kind"] in ("Camera", "Access door")]
            for _ in range(pois(r, s["cameras"] * 0.55)):
                t = self.ts()
                typ = r.choices(types, wts)[0]
                if typ.startswith("Unauthorised deck") and not ship:
                    typ = "Motion in restricted zone after hours"
                _, _, sev, esc = meta[typ]
                an = self.on_shift(t)
                ack = int(r.lognormvariate(3.6 if sev in ("High", "Critical") else 4.3, 0.6))
                target = 60 if sev in ("High", "Critical") else 300
                if typ in ("Camera offline", "Video loss", "NVR disk failure", "Network switch down"):
                    disp = "Technical fault - work order raised" if r.random() < 0.35 else "Self-recovered"
                elif r.random() < esc:
                    disp = "Escalated to incident"
                else:
                    disp = r.choices(DISPOSITIONS, [42, 20, 22, 11, 5])[0]
                aid = f"AL-{self.tag.replace('-', '')}-{len(rows) + 1:06d}"
                cam = r.choice(devs or [{"id": ""}])["id"]
                row = [aid, t.strftime("%Y-%m-%d %H:%M:%S"), self.cl_by[s["client"]]["name"], s["name"], cam,
                       typ, sev, an["name"], ack, target, "Yes" if ack <= target else "No", disp,
                       int(ack + r.lognormvariate(5, 0.8))]
                rows.append(row)
                self.alerts.append((t, s, typ, sev, disp, aid, an))
        rows.sort(key=lambda x: x[1])
        wb = new_book()
        add_sheet(wb, "Alert log", ["Alert ID", "Raised", "Client", "Site", "Device", "Alert type", "Severity",
                                    "Analyst", "Ack seconds", "Ack target s", "Within SLA", "Disposition",
                                    "Close seconds"], rows, {"Client": 30, "Site": 34, "Alert type": 34,
                                                             "Disposition": 34, "Analyst": 22})
        by = {}
        for x in rows:
            k = (x[2], x[6])
            by[k] = by.get(k, 0) + 1
        add_sheet(wb, "By client & severity", ["Client", "Severity", "Alerts"],
                  [[a, b, n] for (a, b), n in sorted(by.items())])
        save_book(self.sink, f"{F['soc']}/{self.y}/SOC alert log {self.tag}.xlsx", wb, "SOC alert log", BANNER)
        # device health
        for s in self.sites:
            for a in self.by_site.get(s["id"], []):
                if a["kind"] not in ("Camera", "NVR"):
                    continue
                off = pois(r, 0.4)
                mins = sum(int(r.expovariate(1 / 50)) for _ in range(off))
                health.append([a["id"], s["name"], a["kind"], a["model"], off, mins,
                               r2(100 - 100 * mins / (self.days * 1440)),
                               r.choice([30, 30, 60, 90]) if a["kind"] == "NVR" else ""])
        wb = new_book()
        add_sheet(wb, "Device uptime", ["Asset ID", "Site", "Kind", "Model", "Offline events", "Offline minutes",
                                        "Uptime %", "Retention days"], health, {"Site": 34, "Model": 32})
        save_book(self.sink, f"{F['soc']}/{self.y}/Device uptime {self.tag}.xlsx", wb, "Device uptime", BANNER)
        # shift roster
        lead = self.people("SOC Shift Lead")
        roster = []
        for d in range(self.days):
            day = self.d0 + dt.timedelta(days=d)
            for si, (sn, h) in enumerate(SHIFTS):
                crew = ((day - dt.date(2016, 1, 1)).days + si) % 4
                an = [e["name"] for e in self.people("SOC Analyst") if e["crew"] == crew]
                ld = [e["name"] for e in lead if e["crew"] == crew] or [lead[0]["name"]]
                roster.append([day.isoformat(), sn, f"{h:02d}:00", f"Crew {'ABCD'[crew]}", ld[0], len(an),
                               ", ".join(an)])
        wb = new_book()
        add_sheet(wb, "Roster", ["Date", "Shift", "Starts", "Crew", "Shift lead", "Analysts", "Names"], roster,
                  {"Names": 90, "Shift lead": 22})
        save_book(self.sink, f"{F['hr']}/{self.y}/SOC shift roster {self.tag}.xlsx", wb, "SOC shift roster", BANNER)

    def incidents(self):
        r = self.rng
        self.incs = []
        cats = {"Door forced open": "Unauthorised entry", "Intrusion - perimeter line crossed": "Intrusion",
                "Smoke / fire analytic": "Fire / life safety", "Abandoned object": "Suspicious object",
                "Failed admin logins on VMS": "Cyber - credential attack",
                "Malware signature on recorder": "Cyber - malware", "Panic button activated": "Personal safety",
                "Unauthorised deck / crew area access": "Unauthorised entry"}
        for t, s, typ, sev, disp, aid, an in self.alerts:
            if disp != "Escalated to incident":
                continue
            pr = {"Critical": "P1", "High": "P2"}.get(sev, "P3")
            if r.random() < 0.15:
                pr = "P1" if pr == "P2" else pr
            cat = cats.get(typ, "Security breach - other")
            close = t + dt.timedelta(hours=r.lognormvariate(1.5 if pr == "P1" else 2.6, 0.7))
            iid = f"INC-{self.tag.replace('-', '')}-{len(self.incs) + 1:04d}"
            loss = r2(r.lognormvariate(7, 1.2)) if r.random() < 0.18 else 0.0
            police = "Yes" if cat in ("Unauthorised entry", "Intrusion", "Personal safety") and r.random() < .35 else "No"
            root = r.choice({
                "Cyber - credential attack": ["Exposed remote-access port", "Default password on recorder",
                                              "Brute force from foreign IP range"],
                "Cyber - malware": ["Unpatched recorder firmware", "Infected USB used for footage export"],
                "Fire / life safety": ["Kitchen extraction fault", "Cigarette in bin", "Analytic false trigger"]
            }.get(cat, ["Tailgating through staff entrance", "Door closer failure", "Contractor without escort",
                        "Guest in restricted area", "Fence damage", "Unknown individual, left before guard arrived"]))
            owner = r.choice(self.people("SOC Shift Lead"))
            self.incs.append({"id": iid, "alert": aid, "opened": t, "closed": close, "site": s, "cat": cat,
                              "pri": pr, "typ": typ, "analyst": an, "owner": owner, "root": root,
                              "police": police, "loss": loss})
        rows = [[i["id"], i["alert"], i["opened"].strftime("%Y-%m-%d %H:%M"), i["closed"].strftime("%Y-%m-%d %H:%M"),
                 self.cl_by[i["site"]["client"]]["name"], i["site"]["name"], i["cat"], i["pri"], i["typ"],
                 i["root"], i["owner"]["name"], i["police"], i["loss"],
                 r2((i["closed"] - i["opened"]).total_seconds() / 3600)] for i in self.incs]
        wb = new_book()
        add_sheet(wb, "Incidents", ["Incident", "Alert", "Opened", "Closed", "Client", "Site", "Category", "Priority",
                                    "Trigger", "Root cause", "Owner", "Police notified", "Loss USD", "Hours open"],
                  rows, {"Client": 28, "Site": 32, "Root cause": 40, "Trigger": 32})
        save_book(self.sink, f"{F['inc']}/{self.y}/Incident register {self.tag}.xlsx", wb, "Incident register", BANNER)
        for i in self.incs:
            if i["pri"] not in ("P1", "P2"):
                continue
            cl = self.cl_by[i["site"]["client"]]
            d = new_doc(f"Incident report {i['id']}", f"{i['cat']} - {i['site']['name']}")
            doc_table(d, ["Field", "Detail"], [
                ["Client", cl["name"]], ["Site", i["site"]["name"]], ["Priority", i["pri"]],
                ["Opened", i["opened"].strftime("%d %b %Y %H:%M")], ["Closed", i["closed"].strftime("%d %b %Y %H:%M")],
                ["Detected by", f"{i['typ']} (alert {i['alert']})"], ["SOC analyst", i["analyst"]["name"]],
                ["Incident owner", i["owner"]["name"]], ["Police notified", i["police"]],
                ["Estimated loss (USD)", f"{i['loss']:,.2f}"]])
            d.add_heading("Timeline", 1)
            t = i["opened"]
            steps = [(0, f"Alert '{i['typ']}' raised by the analytics engine and shown on the SOC video wall."),
                     (r.randint(20, 70) / 60, f"{i['analyst']['name']} acknowledged, reviewed live and recorded footage."),
                     (r.randint(2, 6), f"Client duty manager {cl['contact']} called; on-site security dispatched."),
                     (r.randint(8, 25), "On-site team confirmed the situation and secured the area."),
                     (None, f"Incident closed by {i['owner']['name']} after review.")]
            for mins, text in steps:
                when = i["closed"] if mins is None else t + dt.timedelta(minutes=mins)
                bullet(d, f"{when:%H:%M} - {text}")
            d.add_heading("Root cause", 1)
            d.add_paragraph(f"{i['root']}. Footage from the relevant cameras was preserved for 90 days "
                            "under the evidence-hold procedure.")
            d.add_heading("Corrective actions", 1)
            for a in r.sample(["Re-brief night staff on escort rules", "Adjust analytic zone and sensitivity",
                               "Replace door closer and add held-open alarm", "Rotate recorder credentials",
                               "Patch recorder firmware to current release", "Add camera coverage at blind spot",
                               "Review guard patrol timing", "Restrict remote access to VPN only"], 3):
                bullet(d, a)
            save_doc(self.sink, f"{F['inc']}/{self.y}/Incident reports/{i['id']} {i['cat']}.docx", d)

    def field(self):
        r, rows, parts = self.rng, [], []
        techs = self.people("Field Technician")
        self.billable_parts = {}
        n = 0

        def wo(kind, s, desc, hours, asset="", part=None, cost=0.0, billable=False):
            nonlocal n
            n += 1
            t = self.ts(hour=r.randint(7, 17))
            tech = r.choice(techs)
            wid = f"WO-{self.tag.replace('-', '')}-{n:05d}"
            rows.append([wid, t.strftime("%Y-%m-%d %H:%M"), kind, self.cl_by[s["client"]]["name"], s["name"], asset,
                         desc, tech["name"], r2(hours), part or "", r2(cost), "Yes" if billable else "No",
                         r.choice(["Completed", "Completed", "Completed", "Completed - follow-up needed"])])
            if part:
                parts.append([wid, part, 1, r2(cost)])
            if billable:
                self.billable_parts[s["client"]] = self.billable_parts.get(s["client"], 0) + cost * 1.25
            return wid

        for s in self.sites:
            if s["id"] % 3 == self.ym[1] % 3:
                wo("Preventive maintenance", s, "Quarterly PM: lens cleaning, focus check, NVR health, UPS test",
                   s["cameras"] * r.uniform(0.05, 0.09))
            if s["live"] == self.ym:
                wo("Installation project", s, f"Install and commission {s['cameras']} cameras and recorders",
                   s["cameras"] * r.uniform(2.2, 3.1))
        aid = {a["id"]: a for a in self.w.assets}
        for old, new in self.w.failures.get(self.ym, []):
            a = aid[old]
            s = self.site_by.get(a["site"])
            if not s:
                continue
            age = (self.ym[0] - a["installed"][0]) * 12 + self.ym[1] - a["installed"][1]
            wo("Corrective repair", s, f"Replace failed {a['model']} ({old}) with {new}", r.uniform(1, 4), old,
               a["model"], a["cost"], billable=age > 36)
        self.workorders = rows
        wb = new_book()
        add_sheet(wb, "Work orders", ["Work order", "Scheduled", "Type", "Client", "Site", "Asset", "Description",
                                      "Technician", "Hours", "Part", "Part cost USD", "Billable", "Status"], rows,
                  {"Description": 60, "Site": 32, "Client": 28, "Part": 30})
        add_sheet(wb, "Parts used", ["Work order", "Part", "Qty", "Cost USD"], parts, {"Part": 34})
        save_book(self.sink, f"{F['field']}/{self.y}/Work orders {self.tag}.xlsx", wb, "Work orders", BANNER)
        self.parts_cost = sum(p[3] for p in parts)

    def helpdesk(self):
        r, rows = self.rng, []
        sup = self.people("IT Support Specialist")
        for c in self.clients:
            for _ in range(pois(r, 5 * len(c["sites"]) + 2)):
                cat, pr = r.choice(TICKET_CATS)
                t = self.ts(hour=r.randint(7, 21))
                fr_t, res_t = SLA[pr]
                fr = int(r.lognormvariate(math.log(fr_t * 0.5), 0.7))
                res = r2(r.lognormvariate(math.log(res_t * 0.4), 0.8))
                rows.append([f"HD-{self.tag.replace('-', '')}-{len(rows) + 1:05d}", t.strftime("%Y-%m-%d %H:%M"),
                             c["name"], r.choice(["Email", "Phone", "Portal", "Portal", "WhatsApp"]), cat, pr,
                             r.choice(sup)["name"], fr, fr_t, res, res_t,
                             "Yes" if fr <= fr_t and res <= res_t else "No",
                             r.choices([5, 4, 3, 2, 1, None], [40, 28, 10, 4, 2, 16])[0]])
        rows.sort(key=lambda x: x[1])
        self.tickets = rows
        wb = new_book()
        add_sheet(wb, "Tickets", ["Ticket", "Opened", "Client", "Channel", "Category", "Priority", "Assignee",
                                  "First response min", "Target min", "Resolved hours", "Target hours", "SLA met",
                                  "CSAT"], rows, {"Client": 28, "Category": 30, "Assignee": 22})
        save_book(self.sink, f"{F['help']}/{self.y}/Help desk tickets {self.tag}.xlsx", wb, "Help desk", BANNER)

    def crm(self):
        r, rows = self.rng, []
        am = self.people("Account Manager")
        stages = ["Lead", "Qualified", "Site survey", "Proposal", "Negotiation"]
        for c in self.w.clients:
            lead = (c["since"][0] * 12 + c["since"][1]) - (self.ym[0] * 12 + self.ym[1])
            if 0 <= lead <= 4:
                cams = sum(self.w.sites[s]["cameras"] for s in c["sites"])
                stage = "Closed won" if lead == 0 else stages[max(0, 4 - lead)]
                rows.append([c["name"], c["kind"], c["city"], stage, cams,
                             r2(self.w.contracts[[x["client"] for x in self.w.contracts].index(c["id"])]["monthly"] * 12),
                             {"Lead": 10, "Qualified": 25, "Site survey": 40, "Proposal": 60,
                              "Negotiation": 80, "Closed won": 100}[stage], am[c["id"] % len(am)]["name"]])
        for _ in range(pois(r, 3 + len(self.clients) / 8)):
            kind = r.choice(["Hotel", "Hotel", "Cruise line", "Marina / port", "Residential & retail"])
            nm = f"{r.choice(N.HOTEL_A)} {r.choice(N.HOTEL_B if kind == 'Hotel' else N.CRUISE_B if kind == 'Cruise line' else N.PORT_B if kind == 'Marina / port' else N.TOWER_B)}"
            cams = r.randint(60, 500)
            st = r.choice(stages + ["Closed lost", "Closed lost"])
            rows.append([nm + " (prospect)", kind, N.location(r)[0], st, cams, r2(cams * r.uniform(30, 45) * 12),
                         0 if st == "Closed lost" else r.choice([10, 25, 40, 60]), r.choice(am)["name"]])
        wb = new_book()
        add_sheet(wb, "Pipeline", ["Account", "Segment", "City", "Stage", "Cameras", "Annual value USD",
                                   "Probability %", "Owner"], rows, {"Account": 38, "Owner": 22})
        ws = wb["Pipeline"]
        n = ws.max_row
        ws.cell(n + 2, 5, "Weighted pipeline")
        ws.cell(n + 2, 6, f"=SUMPRODUCT(F2:F{n},G2:G{n})/100")
        save_book(self.sink, f"{F['crm']}/Sales pipeline/{self.y}/Opportunities {self.tag}.xlsx", wb, "Pipeline", BANNER)
        # new contracts as DOCX
        for c in self.w.contracts:
            if c["start"] != self.ym:
                continue
            cl = self.w.clients[c["client"]]
            self.contract_doc(c, cl)

    def contract_doc(self, c, cl):
        d = new_doc(f"Managed Security Services Agreement {c['code']}",
                    f"Between {CO} and {cl['name']} ({c['type'].lower()})")
        d.add_paragraph(f"This Agreement starts on 1 {dt.date(c['start'][0], c['start'][1], 1):%B %Y} and runs for "
                        f"{c['term']} months to the end of {dt.date(c['end'][0], c['end'][1], 1):%B %Y}.")
        d.add_heading("1. Services", 1)
        for s in ["24/7 remote video monitoring from the Security Operations Centre",
                  "Video analytics: intrusion, loitering, abandoned object, fire/smoke detection",
                  "Preventive maintenance each quarter and corrective repair of covered devices",
                  "Cloud video retention (30 days standard, 90 days for evidence holds)",
                  "Help desk for VMS users, footage export and access card administration",
                  "Monthly service report and quarterly service review"]:
            bullet(d, s)
        d.add_heading("2. Covered sites", 1)
        doc_table(d, ["Site", "Cameras"], [[self.w.sites[s]["name"], self.w.sites[s]["cameras"]] for s in cl["sites"]])
        d.add_heading("3. Fees", 1)
        doc_table(d, ["Item", "USD"], [["Base SOC fee per month", c["base_fee"]],
                                       ["Monitoring fee per camera per month", c["fee_per_camera"]],
                                       ["Estimated monthly fee", c["monthly"]],
                                       ["One-off installation fee", c["install_fee"]]])
        d.add_paragraph(f"Invoices are issued monthly in advance and are payable within {cl['terms']} days. "
                        f"Prices exclude VAT / sales tax at {int(cl['vat'] * 100)}% where applicable.")
        d.add_heading("4. Service levels", 1)
        doc_table(d, ["Priority", "Response", "Resolution"], [[p, f"{a} min", f"{b} h"] for p, (a, b) in SLA.items()])
        d.add_paragraph("Critical and high alarms are acknowledged within 60 seconds; other alarms within 5 minutes. "
                        "Service credits of 2% of the monthly fee apply per missed P1 target, capped at 10%.")
        d.add_heading("5. Data protection", 1)
        d.add_paragraph("Recorded video is personal data. The Provider processes it only on the Client's written "
                        "instructions, keeps it encrypted at rest and in transit, and deletes it at the end of the "
                        "retention period unless an evidence hold applies.")
        d.add_heading("6. Term and termination", 1)
        d.add_paragraph("Either party may end the Agreement with 90 days' notice before the renewal date, or "
                        "at once for a material breach not cured within 30 days.")
        d.add_heading("Signatures", 1)
        doc_table(d, ["For", "Name", "Title"], [[CO, self.w.staff[0]["name"], "Chief Executive Officer"],
                                               [cl["name"], cl["contact"], "General Manager"]])
        save_doc(self.sink, f"{F['contracts']}/{safe(cl['name'], 60)}/{c['code']} {c['type']} agreement.docx", d)

    def hr(self):
        r = self.rng
        pay, ts = [], []
        ops = {"SOC Analyst", "SOC Shift Lead", "Field Technician"}
        self.payroll_total = self.eos_total = 0.0
        self.pay_lines = {}
        for e in self.staff:
            base = self.w.pay(e, self.ym)
            house, trans = r2(base * .25), r2(base * .08)
            leave = r.choices([0, 0, 0, 1, 2, 5, 10], [40, 20, 10, 10, 8, 7, 5])[0]
            shifts = (self.days * 3 / 4 * 8 / 3) if e["role"] in ("SOC Analyst", "SOC Shift Lead") else 0
            reg = r2(min(176, (22 - leave) * 8))
            ot = r2(r.choice([0, 0, 4, 8, 12, 16]) if e["role"] in ops else 0)
            otp = r2(ot * base / 176 * 1.25)
            gross = r2(base + house + trans + otp)
            eos = r2(base * 21 / 365 * 12 / 12)
            self.payroll_total += gross
            self.eos_total += eos
            self.pay_lines[e["id"]] = (base, house, trans, ot, otp, gross, eos, leave)
            pay.append([e["id"], e["name"], e["dept"], e["role"], base, house, trans, otp, gross, 0.0, gross, eos,
                        f"AE{r.randrange(10 ** 18):018d}"[:23]])
            ts.append([e["id"], e["name"], e["role"], reg, ot, leave, int(shifts) if shifts else "", self.tag])
        wb = new_book()
        add_sheet(wb, "Payroll", ["Employee", "Name", "Department", "Role", "Basic", "Housing", "Transport",
                                  "Overtime", "Gross", "Deductions", "Net pay", "EOS accrual", "IBAN (fictional)"],
                  pay, {"Name": 22, "Role": 24, "Department": 20, "IBAN (fictional)": 26})
        ws = wb["Payroll"]
        n = ws.max_row
        ws.cell(n + 1, 1, "TOTAL")
        for col in "EFGHIJKL":
            ws[f"{col}{n + 1}"] = f"=SUM({col}2:{col}{n})"
        save_book(self.sink, f"{F['hr']}/{self.y}/Payroll register {self.tag}.xlsx", wb, "Payroll", BANNER)
        wb = new_book()
        add_sheet(wb, "Timesheets", ["Employee", "Name", "Role", "Regular hours", "Overtime hours", "Leave days",
                                     "Shift hours", "Period"], ts, {"Name": 22, "Role": 24})
        save_book(self.sink, f"{F['hr']}/{self.y}/Timesheets {self.tag}.xlsx", wb, "Timesheets", BANNER)
        self.payroll_total, self.eos_total = r2(self.payroll_total), r2(self.eos_total)
        # joiners and leavers
        moves = [[e["id"], e["name"], e["role"], "Joined", self.tag, ""] for e in self.w.staff if e["hired"] == self.ym]
        moves += [[e["id"], e["name"], e["role"], "Left", self.tag, e["leave_reason"]]
                  for e in self.w.staff if e["left"] == self.ym]
        for e in self.w.staff:
            if e["hired"] == self.ym:
                S.joiner_docs(self, e)
            elif e["left"] == self.ym:
                S.leaver_docs(self, e)
        if moves:
            wb = new_book()
            add_sheet(wb, "Movements", ["Employee", "Name", "Role", "Movement", "Month", "Reason"], moves,
                      {"Name": 22, "Role": 24})
            save_book(self.sink, f"{F['hr']}/{self.y}/Joiners and leavers {self.tag}.xlsx", wb, "Movements", BANNER)

    def finance(self):
        r, L = self.rng, self.L
        eom = self.d0 + dt.timedelta(days=self.days - 1)
        S.expenses(self)
        if self.ym == START:
            self.post(self.d0, "OPEN", "Share capital paid in", [("1000", 2500000, 0), ("3000", 0, 2500000)])
        invoices = []
        fin = self.w.contracts
        for c in self.clients:
            con = self.w.contract_for(c["id"], self.ym)
            if not con:
                continue
            lines = [("4000", f"Managed monitoring {self.tag} ({con['cameras']} cameras)", con["monthly"])]
            if con["start"] == self.ym and con["install_fee"]:
                lines.append(("4100", "Installation and commissioning", con["install_fee"]))
            if self.billable_parts.get(c["id"]):
                lines.append(("4200", "Replacement parts outside warranty", r2(self.billable_parts[c["id"]])))
            net = r2(sum(x[2] for x in lines))
            tax = r2(net * c["vat"])
            L["inv_no"] = L.get("inv_no", 0) + 1
            no = f"INV-{L['inv_no']:06d}"
            date = self.d0
            due = date + dt.timedelta(days=c["terms"])
            delay = r.choices([0, 5, 15, 30, 60, 120, None], [30, 25, 20, 12, 8, 4, 1])[0]
            paid = None if delay is None else due + dt.timedelta(days=delay + r.randint(-5, 5))
            invoices.append([no, date.isoformat(), c["code"], c["name"], net, tax, r2(net + tax), due.isoformat()])
            self.post(date, no, f"Invoice {c['name']}", [("1100", net + tax, 0), ("2100", 0, tax)] +
                      [(a, 0, v) for a, _, v in lines])
            L.setdefault("ar", []).append({"no": no, "client": c["name"], "code": c["code"], "date": date.isoformat(),
                                           "due": due.isoformat(), "total": r2(net + tax),
                                           "paid": paid.isoformat() if paid else None})
            self.invoice_pdf(no, date, due, c, lines, net, tax)
            self.mail(self.ts(0, 9), f"accounts@{DOMAIN}", f"payables@{c['domain']}", f"Invoice {no} from {CO}",
                      f"Dear {c['contact']},\n\nPlease find attached invoice {no} for {self.tag}, total USD "
                      f"{net + tax:,.2f}, due {due:%d %b %Y}.\n\nRegards,\nAccounts Receivable\n{CO}\n")
        # receipts
        still = []
        for inv in L.get("ar", []):
            pd = inv["paid"]
            if pd and pd <= eom.isoformat():
                self.post(dt.date.fromisoformat(max(pd, self.d0.isoformat())), "RCPT " + inv["no"],
                          f"Receipt {inv['client']}", [("1000", inv["total"], 0), ("1100", 0, inv["total"])])
            elif not pd and (dt.date.fromisoformat(inv["due"]) + dt.timedelta(days=365)) <= eom:
                self.post(eom, "WO " + inv["no"], f"Bad debt write-off {inv['client']}",
                          [("5900", inv["total"], 0), ("1100", 0, inv["total"])])
            else:
                still.append(inv)
        L["ar"] = still
        # payroll
        self.post(eom, f"PAY-{self.tag}", "Monthly payroll", [("5000", self.payroll_total, 0), ("1000", 0, self.payroll_total)])
        self.post(eom, f"EOS-{self.tag}", "End-of-service accrual", [("5050", self.eos_total, 0), ("2200", 0, self.eos_total)])
        # supplier bills
        cams, nsites, nclients = self.w.load(self.ym)
        staff = len(self.staff)
        bills = []
        costs = {"5100": self.parts_cost, "5200": cams * r.uniform(3.2, 4.1), "5300": nsites * r.uniform(380, 520),
                 "5400": 18000 + staff * 260, "5500": len(self.people("Field Technician")) * 780,
                 "5600": cams * r.uniform(1.4, 1.9) + 2500, "5700": 3200 + cams * 0.35,
                 "6000": staff * r.uniform(40, 90), "6100": 4500 if self.ym[1] in (3, 9) else 900}
        for name, desc, acct in N.SUPPLIERS:
            amt = r2(costs[acct] / (2 if acct == "6000" else 1))
            if amt <= 0:
                continue
            L["bill_no"] = L.get("bill_no", 0) + 1
            no = f"BILL-{L['bill_no']:06d}"
            bdate = self.d0 + dt.timedelta(days=r.randint(0, self.days - 1))
            due = bdate + dt.timedelta(days=30)
            bills.append([no, bdate.isoformat(), name, desc, ACCOUNTS[acct][0], amt, due.isoformat()])
            self.post(bdate, no, f"Bill {name}", [(acct, amt, 0), ("2000", 0, amt)])
            L.setdefault("ap", []).append({"no": no, "supplier": name, "total": amt, "due": due.isoformat()})
        if self.ym[1] in (1, 4, 7, 10):
            capex = r2(60000 + cams * 4)
            L["bill_no"] = L.get("bill_no", 0) + 1
            no = f"BILL-{L['bill_no']:06d}"
            bills.append([no, self.d0.isoformat(), "Northgate Vision Supply", "SOC hardware, vehicles & spares pool",
                          ACCOUNTS["1500"][0], capex, (self.d0 + dt.timedelta(days=30)).isoformat()])
            self.post(self.d0, no, "Capital equipment", [("1500", capex, 0), ("2000", 0, capex)])
            L.setdefault("ap", []).append({"no": no, "supplier": "Northgate Vision Supply", "total": capex,
                                           "due": (self.d0 + dt.timedelta(days=30)).isoformat()})
        still = []
        for b in L.get("ap", []):
            if b["due"] <= eom.isoformat():
                self.post(dt.date.fromisoformat(max(b["due"], self.d0.isoformat())), "PMT " + b["no"],
                          f"Payment {b['supplier']}", [("2000", b["total"], 0), ("1000", 0, b["total"])])
            else:
                still.append(b)
        L["ap"] = still
        # VAT settlement of last month, depreciation
        bal = L.setdefault("bal", {})
        vat_due = r2(-bal.get("2100", 0))
        if vat_due > 0:
            self.post(self.d0 + dt.timedelta(days=27), f"VAT-{self.tag}", "Sales tax return payment",
                      [("2100", vat_due, 0), ("1000", 0, vat_due)])
        dep = r2(bal.get("1500", 0) / 60)
        if dep:
            self.post(eom, f"DEP-{self.tag}", "Monthly depreciation", [("5800", dep, 0), ("1510", 0, dep)])
        for _, _, a, _, _, dr, cr in self.gl:
            bal[a] = r2(bal.get(a, 0) + dr - cr)
        # outputs
        y = self.y
        wb = new_book()
        add_sheet(wb, "Journal", ["Date", "Reference", "Account", "Account name", "Description", "Debit", "Credit"],
                  [[str(g[0])] + g[1:] for g in sorted(self.gl, key=lambda g: str(g[0]))],
                  {"Description": 40, "Account name": 30})
        ws = wb["Journal"]
        n = ws.max_row
        ws.cell(n + 1, 5, "TOTAL")
        ws[f"F{n + 1}"], ws[f"G{n + 1}"] = f"=SUM(F2:F{n})", f"=SUM(G2:G{n})"
        save_book(self.sink, f"{F['fin']}/{y}/General ledger/General ledger {self.tag}.xlsx", wb, "GL", BANNER)
        tb = [[a, ACCOUNTS[a][0], ACCOUNTS[a][1], max(bal.get(a, 0), 0), max(-bal.get(a, 0), 0)]
              for a in sorted(ACCOUNTS)]
        wb = new_book()
        add_sheet(wb, "Trial balance", ["Account", "Name", "Type", "Debit", "Credit"], tb, {"Name": 34})
        ws = wb["Trial balance"]
        n = ws.max_row
        ws.cell(n + 1, 2, "TOTAL")
        ws[f"D{n + 1}"], ws[f"E{n + 1}"] = f"=SUM(D2:D{n})", f"=SUM(E2:E{n})"
        ws.cell(n + 2, 2, "Difference")
        ws[f"D{n + 2}"] = f"=D{n + 1}-E{n + 1}"
        save_book(self.sink, f"{F['fin']}/{y}/Trial balance/Trial balance {self.tag}.xlsx", wb, "TB", BANNER)
        wb = new_book()
        add_sheet(wb, "Invoices", ["Invoice", "Date", "Client code", "Client", "Net USD", "Tax USD", "Total USD", "Due"],
                  invoices, {"Client": 30})
        add_sheet(wb, "Supplier bills", ["Bill", "Date", "Supplier", "Description", "Expense account", "Amount USD", "Due"],
                  bills, {"Supplier": 28, "Description": 34, "Expense account": 28})
        save_book(self.sink, f"{F['fin']}/{y}/Sales and purchases/Invoices and bills {self.tag}.xlsx", wb, "AR/AP", BANNER)
        aging = []
        for inv in L["ar"]:
            days = (eom - dt.date.fromisoformat(inv["due"])).days
            b = "Current" if days <= 0 else "1-30" if days <= 30 else "31-60" if days <= 60 else "61-90" \
                if days <= 90 else "90+"
            aging.append([inv["no"], inv["client"], inv["date"], inv["due"], inv["total"], max(days, 0), b])
        wb = new_book()
        add_sheet(wb, "AR aging", ["Invoice", "Client", "Date", "Due", "Open USD", "Days overdue", "Bucket"], aging,
                  {"Client": 30})
        save_book(self.sink, f"{F['fin']}/{y}/Receivables/AR aging {self.tag}.xlsx", wb, "AR aging", BANNER)
        bank, run = [], L.get("cash_open", 0.0)
        for g in sorted([g for g in self.gl if g[2] == "1000"], key=lambda g: str(g[0])):
            run = r2(run + g[5] - g[6])
            bank.append([str(g[0]), g[1], g[4], g[5] or "", g[6] or "", run])
        L["cash_open"] = run
        buf = io.StringIO()
        buf.write("Date,Reference,Description,Credit (in),Debit (out),Balance\n")
        for b in bank:
            buf.write(",".join(f'"{x}"' if isinstance(x, str) else str(x) for x in b) + "\n")
        self.sink.write_text(f"{F['fin']}/{y}/Bank statements/Bank statement {self.tag}.csv", buf.getvalue())
        pl = {a: r2(-sum(dr - cr for _, _, aa, _, _, dr, cr in self.gl if aa == a)) for a in ACCOUNTS
              if ACCOUNTS[a][1] in ("Revenue", "Expense")}
        L.setdefault("pl", {})[self.tag] = pl
        self.pl = pl
        if self.ym[1] == 12:
            self.year_end(eom)
        if self.ym[1] == 1:
            self.budget()

    def invoice_pdf(self, no, date, due, c, lines, net, tax):
        p = self.sink.path(f"{F['fin']}/{self.y}/Invoices/{self.tag}/{no} {safe(c['name'], 50)}.pdf")
        st = getSampleStyleSheet()
        doc = SimpleDocTemplate(p, pagesize=A4, title=f"Invoice {no}", author=CO, subject=BANNER)
        rows = [["Description", "Account", "Amount USD"]] + [[d, a, f"{v:,.2f}"] for a, d, v in lines] + \
               [["", "Net", f"{net:,.2f}"], ["", f"Tax {int(c['vat'] * 100)}%", f"{tax:,.2f}"],
                ["", "Total due", f"{net + tax:,.2f}"]]
        t = Table(rows, colWidths=[300, 80, 100])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("ALIGN", (2, 0), (2, -1), "RIGHT"),
                               ("GRID", (0, 0), (-1, len(lines)), 0.4, colors.grey),
                               ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold")]))
        doc.build([Paragraph(f"<b>{CO}</b>", st["Title"]), Paragraph(f"TAX INVOICE {no}", st["Heading2"]),
                   Paragraph(f"Bill to: {c['name']}, {c['city']}, {c['country']}<br/>Attn: {c['contact']}<br/>"
                             f"Invoice date: {date:%d %b %Y} &nbsp; Due: {due:%d %b %Y} &nbsp; Client: {c['code']}",
                             st["Normal"]), Spacer(1, 14), t, Spacer(1, 18),
                   Paragraph(f"Pay by bank transfer quoting {no}. {BANNER}", st["Italic"])])
        self.sink.wrote(p)

    def year_end(self, eom):
        L, y = self.L, self.ym[0]
        months = [f"{y}-{m:02d}" for m in range(1, 13) if f"{y}-{m:02d}" in L["pl"]]
        bal = L["bal"]
        wb = new_book()
        ws = wb.create_sheet("Income statement")
        ws.append([f"{CO} - Income statement {y} (USD, fictional)"])
        ws.append(["Account", "Name"] + months + ["Full year"])
        rev = [a for a in sorted(ACCOUNTS) if ACCOUNTS[a][1] == "Revenue"]
        exp = [a for a in sorted(ACCOUNTS) if ACCOUNTS[a][1] == "Expense"]
        last = chr(ord("C") + len(months) - 1)
        tot = chr(ord(last) + 1)
        for a in rev:
            ws.append([a, ACCOUNTS[a][0]] + [L["pl"][m].get(a, 0) for m in months])
            ws[f"{tot}{ws.max_row}"] = f"=SUM(C{ws.max_row}:{last}{ws.max_row})"
        r0, r1 = 3, ws.max_row
        ws.append(["", "Total revenue"] + [f"=SUM({chr(67 + i)}{r0}:{chr(67 + i)}{r1})" for i in range(len(months) + 1)])
        rr = ws.max_row
        for a in exp:
            ws.append([a, ACCOUNTS[a][0]] + [-L["pl"][m].get(a, 0) for m in months])
            ws[f"{tot}{ws.max_row}"] = f"=SUM(C{ws.max_row}:{last}{ws.max_row})"
        e0, e1 = rr + 1, ws.max_row
        ws.append(["", "Total expenses"] + [f"=SUM({chr(67 + i)}{e0}:{chr(67 + i)}{e1})" for i in range(len(months) + 1)])
        re_ = ws.max_row
        ws.append(["", "Net profit"] + [f"={chr(67 + i)}{rr}-{chr(67 + i)}{re_}" for i in range(len(months) + 1)])
        ws.column_dimensions["B"].width = 36
        bs = wb.create_sheet("Balance sheet")
        bs.append([f"{CO} - Balance sheet at 31 Dec {y} (USD, fictional)"])
        bs.append(["Account", "Name", "Amount"])
        for kind, sign in (("Asset", 1), ("Liability", -1), ("Equity", -1)):
            bs.append(["", kind.upper() + ("S" if kind != "Equity" else "")])
            s0 = bs.max_row + 1
            for a in sorted(ACCOUNTS):
                if ACCOUNTS[a][1] == kind:
                    bs.append([a, ACCOUNTS[a][0], r2(sign * bal.get(a, 0))])
            if kind == "Equity":
                bs.append(["", "Current year profit", f"='Income statement'!{tot}{re_ + 1}"])
            bs.append(["", f"Total {kind.lower()}", f"=SUM(C{s0}:C{bs.max_row})"])
        bs.append(["", "Check: assets - liabilities - equity", "=C{0}-C{1}-C{2}".format(
            *[i for i in range(1, bs.max_row + 1) if str(bs.cell(i, 2).value).startswith("Total ")])])
        bs.column_dimensions["B"].width = 40
        save_book(self.sink, f"{F['fin']}/{y}/Financial statements {y}.xlsx", wb, "Financial statements", BANNER)
        # closing entry: move P&L to retained earnings
        lines, net = [], 0.0
        for a in rev + exp:
            v = bal.get(a, 0)
            if v:
                lines.append((a, max(-v, 0), max(v, 0)))
                net += v
        if lines:
            lines.append(("3100", max(net, 0), max(-net, 0)))
            for a, dr, cr in lines:
                bal[a] = r2(bal.get(a, 0) + dr - cr)
            buf = [[eom.isoformat(), f"CLOSE-{y}", a, ACCOUNTS[a][0], "Year-end close to retained earnings", r2(dr), r2(cr)]
                   for a, dr, cr in lines if dr or cr]
            wb = new_book()
            add_sheet(wb, "Closing entries", ["Date", "Reference", "Account", "Account name", "Description", "Debit",
                                              "Credit"], buf, {"Account name": 30, "Description": 36})
            save_book(self.sink, f"{F['fin']}/{y}/General ledger/Year-end closing entries {y}.xlsx", wb, "Close", BANNER)

    def budget(self):
        L, y = self.L, self.ym[0]
        prev = [f"{y - 1}-{m:02d}" for m in range(1, 13) if f"{y - 1}-{m:02d}" in L.get("pl", {})]
        if not prev:
            return
        wb = new_book()
        ws = wb.create_sheet("Budget")
        ws.append([f"{CO} - Operating budget {y} (USD, fictional)"])
        ws.append(["Growth assumption", 0.12])
        ws.append(["Account", "Name", f"{y - 1} actual"] + [dt.date(y, m, 1).strftime("%b") for m in range(1, 13)] + [f"{y} total"])
        for a in sorted(ACCOUNTS):
            if ACCOUNTS[a][1] not in ("Revenue", "Expense"):
                continue
            act = r2(sum(abs(L["pl"][m].get(a, 0)) for m in prev))
            row = ws.max_row + 1
            ws.append([a, ACCOUNTS[a][0], act] + [f"=$C{row}*(1+$B$2)/12" for _ in range(12)] + [f"=SUM(D{row}:O{row})"])
        ws.column_dimensions["B"].width = 34
        save_book(self.sink, f"{F['fin']}/{y}/Budget {y}.xlsx", wb, "Budget", BANNER)

    def reports(self):
        r = self.rng
        kpi = []
        for c in self.clients:
            sids = set(c["sites"])
            al = [a for a in self.alerts if a[1]["id"] in sids]
            inc = [i for i in self.incs if i["site"]["id"] in sids]
            tk = [t for t in self.tickets if t[2] == c["name"]]
            wos = [w for w in self.workorders if w[3] == c["name"]]
            sla_al = sum(1 for a in al if a[4] != "") and r2(100 * sum(
                1 for a in al if True) / max(len(al), 1))
            met = r2(100 * sum(1 for t in tk if t[11] == "Yes") / max(len(tk), 1))
            fp = r2(100 * sum(1 for a in al if a[4].startswith("False")) / max(len(al), 1))
            cams = sum(self.w.sites[s]["cameras"] for s in c["sites"])
            kpi.append([c["code"], c["name"], cams, len(al), fp, len(inc), sum(1 for i in inc if i["pri"] == "P1"),
                        len(tk), met, len(wos)])
            d = new_doc(f"Monthly service report - {c['name']}", f"{dt.date(*self.ym, 1):%B %Y}")
            d.add_paragraph(f"Prepared for {c['contact']}. Coverage: {cams} cameras across {len(sids)} site(s).")
            d.add_heading("Summary", 1)
            doc_table(d, ["Measure", "Value"], [["Alarms handled", len(al)], ["False alarm rate %", fp],
                                                ["Incidents", len(inc)], ["Help desk tickets", len(tk)],
                                                ["Tickets within SLA %", met], ["Work orders", len(wos)]])
            by = {}
            for a in al:
                by[a[2]] = by.get(a[2], 0) + 1
            d.add_heading("Alarms by type", 1)
            doc_table(d, ["Alarm type", "Count"], sorted(by.items(), key=lambda x: -x[1])[:12])
            if inc:
                d.add_heading("Incidents", 1)
                doc_table(d, ["Incident", "Priority", "Category", "Site", "Root cause"],
                          [[i["id"], i["pri"], i["cat"], i["site"]["name"], i["root"]] for i in inc])
            d.add_heading("Recommendations", 1)
            recs = []
            if fp > 60:
                recs.append("Tune analytic zones on the cameras with the most false alarms; we will propose "
                            "schedules that match guest and staff traffic.")
            if any(i["cat"].startswith("Cyber") for i in inc):
                recs.append("Move remote access for recorders behind the VPN and rotate all admin credentials.")
            if any(i["cat"] == "Unauthorised entry" for i in inc):
                recs.append("Refresh staff-entrance tailgating briefing and add a held-open alarm on back doors.")
            recs += r.sample(["Plan replacement of cameras older than five years in the next budget.",
                              "Add cloud retention of 60 days for the main lobby and cash office cameras.",
                              "Schedule a joint drill with the on-site security team next quarter.",
                              "Review user accounts on the VMS; remove leavers each month."], 2)
            for x in recs:
                bullet(d, x)
            save_doc(self.sink, f"{F['rep']}/{self.y}/Client service reports/{self.tag}/{safe(c['name'], 60)} {self.tag}.docx", d)
            if inc:
                i = inc[0]
                self.mail(i["opened"] + dt.timedelta(minutes=12), f"soc@{DOMAIN}", f"security@{c['domain']}",
                          f"[{i['pri']}] {i['cat']} at {i['site']['name']} - {i['id']}",
                          f"Dear {c['contact']},\n\nThe SOC raised incident {i['id']} ({i['cat']}) at "
                          f"{i['opened']:%H:%M} on {i['opened']:%d %b}. Trigger: {i['typ']}. On-site security has "
                          f"been informed and footage is preserved.\n\nSOC duty lead: {i['owner']['name']}\n{CO} SOC\n")
        wb = new_book()
        add_sheet(wb, "Client KPIs", ["Code", "Client", "Cameras", "Alarms", "False alarm %", "Incidents", "P1",
                                      "Tickets", "Tickets in SLA %", "Work orders"], kpi, {"Client": 32})
        ws = wb["Client KPIs"]
        n = ws.max_row
        ws.cell(n + 1, 2, "TOTAL")
        for col in "CDFGHJ":
            ws[f"{col}{n + 1}"] = f"=SUM({col}2:{col}{n})"
        fin = wb.create_sheet("Finance")
        fin.append(["Account", "Name", "Amount (USD, credit positive)"])
        for a, v in sorted(self.pl.items()):
            fin.append([a, ACCOUNTS[a][0], v])
        fin.append(["", "Net profit", f"=SUM(C2:C{fin.max_row})"])
        fin.column_dimensions["B"].width = 34
        save_book(self.sink, f"{F['rep']}/{self.y}/KPI dashboards/KPI dashboard {self.tag}.xlsx", wb, "KPIs", BANNER)
        if self.ym[1] % 3 == 0:
            self.board_pack()

    def board_pack(self):
        L, q = self.L, (self.ym[1] - 1) // 3 + 1
        ms = [ym_s(ym_add(self.ym, -k)) for k in (2, 1, 0)]
        pl = [L["pl"].get(m, {}) for m in ms]
        rev = [sum(v for a, v in p.items() if ACCOUNTS[a][1] == "Revenue") for p in pl]
        exp = [-sum(v for a, v in p.items() if ACCOUNTS[a][1] == "Expense") for p in pl]
        cams, sites, clients = self.w.load(self.ym)
        d = new_doc(f"Board pack Q{q} {self.y}", f"{CO} - quarter ending {dt.date(*self.ym, 1):%B %Y}")
        d.add_heading("1. Financial performance", 1)
        doc_table(d, ["Month", "Revenue USD", "Expenses USD", "Net USD"],
                  [[m, r2(a), r2(b), r2(a - b)] for m, a, b in zip(ms, rev, exp)] +
                  [["Quarter", r2(sum(rev)), r2(sum(exp)), r2(sum(rev) - sum(exp))]])
        d.add_paragraph(f"Cash at bank at quarter end: USD {L['bal'].get('1000', 0):,.2f}. "
                        f"Open receivables: USD {sum(x['total'] for x in L['ar']):,.2f} across {len(L['ar'])} invoices.")
        d.add_heading("2. Operations", 1)
        doc_table(d, ["Measure", "Value"], [["Active clients", clients], ["Sites monitored", sites],
                                            ["Cameras monitored", cams], ["Headcount", len(self.staff)],
                                            ["Alarms this month", len(self.alerts)], ["Incidents this month", len(self.incs)]])
        d.add_heading("3. Commercial", 1)
        new = [c["name"] for c in self.w.clients if ym_add(self.ym, -2) <= ym_add(c["since"], 1) <= self.ym]
        lost = [c["name"] for c in self.w.clients if c["end"] and ym_add(self.ym, -2) <= c["end"] <= self.ym]
        d.add_paragraph("New clients: " + (", ".join(new) or "none") + ".")
        d.add_paragraph("Clients lost: " + (", ".join(lost) or "none") + ".")
        d.add_heading("4. Risks", 1)
        for x in self.rng.sample(["Recorder firmware lag on older sites (cyber exposure).",
                                  "Analyst attrition above plan on night crews.",
                                  "Concentration: top five clients exceed 40% of revenue.",
                                  "Supplier lead times for PTZ cameras at 10+ weeks.",
                                  "Data-protection rules for cross-border video storage are tightening."], 3):
            bullet(d, x)
        save_doc(self.sink, f"{F['rep']}/{self.y}/Board packs/Board pack {self.y} Q{q}.docx", d)

    def meetings(self):
        r = self.rng
        mgrs = [e for e in self.staff if e["role"] in HEADS.values()]
        for wk, day in enumerate(range(0, self.days, 7)):
            date = self.d0 + dt.timedelta(days=day)
            al = [a for a in self.alerts if (a[0].date() - date).days in range(-7, 0)]
            inc = [i for i in self.incs if (i["opened"].date() - date).days in range(-7, 0)]
            d = new_doc(f"Weekly operations meeting - {date:%d %b %Y}")
            d.add_paragraph("Attendees: " + ", ".join(f"{e['name']} ({e['role']})" for e in mgrs))
            d.add_heading("1. Last week in numbers", 1)
            bullet(d, f"{len(al)} alarms handled; {sum(1 for a in al if a[4].startswith('False'))} false alarms.")
            bullet(d, f"{len(inc)} incidents, {sum(1 for i in inc if i['pri'] == 'P1')} at P1.")
            bullet(d, f"Headcount {len(self.staff)}; open SOC analyst roles: {r.randint(0, 3)}.")
            d.add_heading("2. Discussion", 1)
            for x in r.sample(["Night-shift coverage during the holiday period", "Recorder firmware rollout status",
                               "Spare-parts stock for 4K bullet cameras", "Client escalations and service credits",
                               "Cloud storage costs per camera", "New client onboarding plan",
                               "Overdue receivables over 60 days", "Technician van utilisation"], 3):
                bullet(d, x)
            d.add_heading("3. Actions", 1)
            doc_table(d, ["Action", "Owner", "Due"], [[x, r.choice(mgrs)["name"], (date + dt.timedelta(days=r.randint(3, 14))).isoformat()]
                                                      for x in r.sample(["Publish revised escalation matrix",
                                                                         "Chase overdue invoices",
                                                                         "Book recorder firmware window",
                                                                         "Interview analyst candidates",
                                                                         "Review false-alarm top 20 cameras",
                                                                         "Send renewal proposals"], 3)])
            save_doc(self.sink, f"{F['mail']}/{self.y}/Meeting minutes/{date} Weekly operations meeting.docx", d)
            self.mail(dt.datetime.combine(date, dt.time(8, 30)), f"ops@{DOMAIN}", f"managers@{DOMAIN}",
                      f"Weekly ops meeting {date:%d %b} - agenda",
                      "Team,\n\nAgenda for today: last week's numbers, firmware rollout, staffing, receivables.\n\nOps\n")

    def write_mail(self):
        for wo in self.workorders[:60]:
            c = next((x for x in self.clients if x["name"] == wo[3]), None)
            if c:
                t = dt.datetime.strptime(wo[1], "%Y-%m-%d %H:%M") + dt.timedelta(hours=wo[8] + 1)
                self.mail(t, f"field@{DOMAIN}", f"facilities@{c['domain']}", f"Work order {wo[0]} completed",
                          f"Hello,\n\n{wo[2]} at {wo[4]} is complete: {wo[6]}. Technician: {wo[7]}.\n\n{CO} Field Services\n")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for i, (t, subj, raw) in enumerate(sorted(self.emails, key=lambda x: x[0]), 1):
                z.writestr(f"{i:05d} {safe(subj, 70)}.eml", raw)
        self.sink.write_bytes(f"{F['mail']}/{self.y}/Email exports/Mailbox export {self.tag}.zip", buf.getvalue())

    # -------- machine telemetry (bulk gzip CSV)
    def _gz(self, rel, header):
        p = self.sink.path(rel)
        f = gzip.open(p, "wt", encoding="utf-8", compresslevel=6, newline="")
        f.write(",".join(header) + "\n")
        return p, f

    def telemetry(self):
        r, k = self.rng, self.detail
        base = f"{F['soc']}/{self.y}/Telemetry {self.tag}"
        t0 = dt.datetime.combine(self.d0, dt.time())
        secs = self.days * 86400
        badges = [e["id"] for e in self.staff]
        # access control door events
        p, f = self._gz(f"{base}/Access control events {self.tag}.csv.gz",
                        ["timestamp", "site", "door", "credential", "holder_type", "event", "direction"])
        evs = ["Access granted"] * 80 + ["Access denied - invalid credential"] * 4 + ["Access denied - schedule"] * 3 \
            + ["Exit request (REX)"] * 9 + ["Door held open"] * 3 + ["Door forced open"]
        for s in self.sites:
            doors = [a["id"] for a in self.by_site.get(s["id"], []) if a["kind"] == "Access door"]
            for door in doors:
                n = pois(r, 22 * self.days * k)
                ts = sorted(r.randrange(secs) for _ in range(n))
                for x in ts:
                    ev = r.choice(evs)
                    staff = r.random() < 0.35
                    cred = r.choice(badges) if staff else f"G{r.randrange(16 ** 7):07X}"
                    f.write(f"{(t0 + dt.timedelta(seconds=x)):%Y-%m-%d %H:%M:%S},{s['code']},{door},{cred},"
                            f"{'Staff' if staff else r.choice(('Guest', 'Guest', 'Contractor', 'Crew'))},{ev},"
                            f"{'Out' if ev.startswith('Exit') else 'In'}\n")
        f.close()
        self.sink.wrote(p)
        # recorder health every 5 minutes (thinned by detail)
        step = max(300, int(300 / max(k, 0.01)))
        p, f = self._gz(f"{base}/Recorder health metrics {self.tag}.csv.gz",
                        ["timestamp", "site", "recorder", "cpu_pct", "mem_pct", "disk_pct", "ingest_mbps",
                         "channels_online", "temp_c"])
        for s in self.sites:
            for a in self.by_site.get(s["id"], []):
                if a["kind"] != "NVR":
                    continue
                disk, chans = r.uniform(55, 80), 48
                for x in range(0, secs, step):
                    hr = (x // 3600) % 24
                    load = 0.75 + 0.25 * math.sin((hr - 6) / 24 * 2 * math.pi)
                    disk = min(97, max(40, disk + r.gauss(0, 0.05)))
                    on = chans - (r.random() < 0.01) * r.randint(1, 4)
                    f.write(f"{(t0 + dt.timedelta(seconds=x)):%Y-%m-%d %H:%M},{s['code']},{a['id']},"
                            f"{min(99, 30 * load + r.gauss(0, 4)):.1f},{55 + r.gauss(0, 3):.1f},{disk:.1f},"
                            f"{on * 4.1 * load + r.gauss(0, 3):.1f},{on},{41 + 6 * load + r.gauss(0, 1):.1f}\n")
        f.close()
        self.sink.wrote(p)
        # device syslog
        msgs = [("info", "vms", "User {u} logged in from 10.{a}.{b}.{c}"), ("info", "vms", "Export job {n} completed"),
                ("warning", "nvr", "Channel {ch} bitrate below threshold"), ("info", "nvr", "Recording schedule applied"),
                ("error", "nvr", "Channel {ch} video loss"), ("info", "nvr", "Channel {ch} video restored"),
                ("warning", "switch", "Port {ch} PoE budget at {n}%"), ("info", "switch", "Port {ch} link up 1000Mbps"),
                ("warning", "switch", "Port {ch} link down"), ("notice", "vms", "Analytics rule {n} triggered"),
                ("warning", "vms", "Failed login for user admin from 203.0.113.{c}"),
                ("info", "acs", "Controller heartbeat ok"), ("error", "acs", "Door {ch} sensor fault"),
                ("info", "ntp", "Clock synchronised, offset {n}ms"), ("notice", "fw", "Firmware check: current")]
        p, f = self._gz(f"{base}/Device syslog {self.tag}.csv.gz", ["timestamp", "site", "host", "severity",
                                                                    "facility", "message"])
        for s in self.sites:
            hosts = [a for a in self.by_site.get(s["id"], []) if a["kind"] in ("NVR", "Switch", "Server")]
            if not hosts:
                continue
            n = pois(r, 9 * len(hosts) * self.days * k)
            for x in sorted(r.randrange(secs) for _ in range(n)):
                sev, fac, m = r.choice(msgs)
                h = r.choice(hosts)
                m = m.format(u=r.choice(badges), a=r.randint(1, 250), b=r.randint(0, 255), c=r.randint(2, 254),
                             n=r.randint(1, 99), ch=r.randint(1, 48))
                f.write(f"{(t0 + dt.timedelta(seconds=x)):%Y-%m-%dT%H:%M:%S}Z,{s['code']},{h['kind'].lower()}-"
                        f"{h['id'][-5:]},{sev},{fac},\"{m}\"\n")
        f.close()
        self.sink.wrote(p)


# ---------------------------------------------------------------- snapshots & static documents

def snapshots(w, ym, sink):
    rows = []
    for c in w.clients:
        if c["since"] >= ym:
            continue
        con = w.contract_for(c["id"], ym)
        cams = sum(w.sites[s]["cameras"] for s in c["sites"])
        rows.append([c["code"], c["name"], c["kind"], c["city"], c["country"], ym_s(ym_add(c["since"], 1)),
                     ym_s(c["end"]) if c["end"] and c["end"] <= ym else "", "Lost" if c["end"] and c["end"] <= ym else "Active",
                     len(c["sites"]), cams, con["code"] if con else "", con["monthly"] if con else "", c["terms"],
                     c["contact"], f"security@{c['domain']}"])
    wb = new_book()
    add_sheet(wb, "Clients", ["Code", "Client", "Segment", "City", "Country", "Client since", "Ended", "Status", "Sites",
                              "Cameras", "Current contract", "Monthly fee USD", "Payment terms", "Contact",
                              "Contact email"], rows, {"Client": 32, "Contact": 22, "Contact email": 30})
    add_sheet(wb, "Sites", ["Site", "Name", "Client", "Cameras", "Live from"],
              [[s["code"], s["name"], w.clients[s["client"]]["name"], s["cameras"], ym_s(s["live"])]
               for s in w.sites if s["live"] <= ym], {"Name": 34, "Client": 30})
    add_sheet(wb, "Contracts", ["Contract", "Client", "Type", "Start", "End", "Term", "Cameras", "Per camera",
                                "Base fee", "Monthly", "Install fee", "Status"],
              [[c["code"], w.clients[c["client"]]["name"], c["type"], ym_s(c["start"]), ym_s(c["end"]), c["term"],
                c["cameras"], c["fee_per_camera"], c["base_fee"], c["monthly"], c["install_fee"],
                c["status"] if c["end"] < ym else ("Active" if c["start"] <= ym else "Signed")]
               for c in w.contracts if c["start"] <= ym_add(ym, 1)], {"Client": 30})
    save_book(sink, f"{F['crm']}/Client master (as of {ym_s(ym)}).xlsx", wb, "Client master", BANNER)
    assets = [[a["id"], a["kind"], a["model"], a["serial"], a["ip"], w.sites[a["site"]]["name"],
               w.clients[w.sites[a["site"]]["client"]]["name"], ym_s(a["installed"]),
               ym_s(ym_add(a["installed"], 36)), a["cost"],
               a["status"] if a["retired"] and a["retired"] <= ym else "In service",
               ym_s(a["retired"]) if a["retired"] and a["retired"] <= ym else ""]
              for a in w.assets if a["installed"] <= ym]
    wb = new_book()
    add_sheet(wb, "Assets", ["Asset ID", "Kind", "Model", "Serial", "IP address", "Site", "Client", "Installed",
                             "Warranty to", "Unit cost USD", "Status", "Retired"], assets,
              {"Model": 32, "Site": 32, "Client": 30, "Status": 24})
    save_book(sink, f"{F['assets']}/Asset register (as of {ym_s(ym)}).xlsx", wb, "Asset register", BANNER)
    staff = [[e["id"], e["name"], e["email"], e["role"], e["dept"], e["manager"], e["nationality"],
              ym_s(e["hired"]), ym_s(e["left"]) if e["left"] and e["left"] <= ym else "", e["leave_reason"]
              if e["left"] and e["left"] <= ym else "", w.pay(e, ym), "ABCD"[e["crew"]] if "SOC" in e["role"] else ""]
             for e in w.staff if e["hired"] <= ym]
    wb = new_book()
    add_sheet(wb, "Employees", ["Employee", "Name", "Email", "Role", "Department", "Manager", "Nationality", "Hired",
                                "Left", "Leave reason", "Basic USD", "SOC crew"], staff,
              {"Name": 22, "Email": 34, "Role": 24, "Department": 20})
    save_book(sink, f"{F['hr']}/Employee roster (as of {ym_s(ym)}).xlsx", wb, "Employee roster", BANNER)


POLICIES = {
    "Information Security Policy": ["Purpose and scope", "Roles and responsibilities", "Asset management",
                                    "Access control", "Cryptography", "Operations security", "Supplier security",
                                    "Incident management", "Compliance and review"],
    "CCTV Monitoring Standard Operating Procedure": ["Video wall layout", "Alarm triage", "Verification steps",
                                                     "Client notification", "Guard dispatch", "Evidence preservation",
                                                     "Shift handover", "Quality checks"],
    "Incident Response Plan": ["Severity definitions", "Escalation matrix", "Containment", "Communication",
                               "Evidence handling", "Recovery", "Post-incident review"],
    "Video Data Retention and Privacy Policy": ["Lawful basis", "Retention periods", "Access requests",
                                                "Footage export to police", "Deletion", "Cross-border storage"],
    "Field Maintenance Procedure": ["Preventive schedule", "Work order lifecycle", "Spare parts", "Safety at height",
                                    "Working on vessels", "Commissioning checklist"],
    "Business Continuity Plan": ["Critical services", "SOC failover site", "Power and connectivity",
                                 "Staff availability", "Testing"],
    "Acceptable Use Policy": ["Company devices", "Email and messaging", "Social media", "Monitoring",
                              "Prohibited activity"],
    "Vendor Management Policy": ["Due diligence", "Contract clauses", "Performance review", "Offboarding"],
    "Health and Safety Policy": ["Risk assessment", "PPE", "Lone working", "Driving", "Reporting accidents"],
    "Remote Access and Recorder Hardening Standard": ["Network segmentation", "Credentials", "Firmware",
                                                      "Logging", "Vulnerability management"],
}
POLICY_TEXT = [
    "{co} applies this section to every employee, contractor and site it serves.",
    "Managers check compliance each quarter and report exceptions to the Operations Director.",
    "Records that show this control works are kept for at least three years.",
    "Exceptions need written approval from the process owner and expire after twelve months.",
    "Training on this topic is part of induction and is refreshed every year.",
    "Where a client's own rules are stricter, the client's rules apply on that client's sites.",
]


def static_docs(w, sink):
    r = random.Random(SEED + 7)
    for i, (title, secs) in enumerate(POLICIES.items(), 1):
        d = new_doc(title, f"Document HL-POL-{i:02d} | Version {r.randint(2, 6)}.{r.randint(0, 9)}")
        doc_table(d, ["Owner", "Approved by", "Review cycle"],
                  [[w.staff[3]["name"] if len(w.staff) > 3 else "SOC Manager", w.staff[0]["name"], "Annual"]])
        for j, s in enumerate(secs, 1):
            d.add_heading(f"{j}. {s}", 1)
            for t in r.sample(POLICY_TEXT, 3):
                d.add_paragraph(t.format(co=CO))
        save_doc(sink, f"{F['pol']}/{title}.docx", d)
    for role, (dept, base, _) in ROLES.items():
        d = new_doc(f"Job description - {role}", f"Department: {dept}")
        doc_table(d, ["Reports to", "Grade salary (USD/month)"], [[HEADS.get(dept, "Chief Executive Officer"),
                                                                  f"{base * .9:,.0f} - {base * 1.15:,.0f}"]])
        d.add_heading("Purpose", 1)
        d.add_paragraph(f"The {role} supports {CO}'s managed security and technology services for hotels, cruise "
                        "lines, marinas and residential clients.")
        d.add_heading("Key responsibilities", 1)
        for x in r.sample(["Meet the service levels in client contracts", "Keep accurate records in the ticketing system",
                           "Follow the information security and privacy policies", "Escalate risks early",
                           "Support continuous improvement of procedures", "Coach new team members",
                           "Take part in incident reviews"], 4):
            bullet(d, x)
        save_doc(sink, f"{F['hr']}/Job descriptions/{role}.docx", d)
    S.it_sops(w, sink, sys.modules[__name__])
    d = new_doc(f"{CO} - dataset README", BANNER)
    d.add_paragraph("This folder holds a synthetic but internally consistent set of operating records for a "
                    "fictional managed security and IT services company. Invoices tie to the general ledger, "
                    "the trial balance balances every month, work orders refer to assets in the register, "
                    "incidents refer to SOC alerts, and payroll covers everyone on the roster for that month.")
    for k, v in F.items():
        bullet(d, v.split("/", 1)[1])
    save_doc(sink, f"{ROOT}/README.docx", d)
    dd = [
        ["01 Clients & CRM", "Client master (as of …).xlsx", "Clients, sites, contracts snapshot"],
        ["01 Clients & CRM", "Sales pipeline/YYYY/Opportunities YYYY-MM.xlsx", "Opportunities and weighted pipeline"],
        ["02 Contracts", "<client>/MSA-… agreement.docx", "New and renewal service agreements"],
        ["03 SOC Monitoring", "YYYY/SOC alert log YYYY-MM.xlsx", "Every alarm: device, analyst, ack time, disposition"],
        ["03 SOC Monitoring", "YYYY/Device uptime YYYY-MM.xlsx", "Camera and recorder uptime"],
        ["04 Incidents", "YYYY/Incident register YYYY-MM.xlsx + Incident reports/*.docx", "Escalated alarms"],
        ["05 Field Service & Maintenance", "YYYY/Work orders YYYY-MM.xlsx", "PM, installs, repairs, parts"],
        ["06 IT Help Desk", "YYYY/Help desk tickets YYYY-MM.xlsx", "Client tickets with SLA and CSAT"],
        ["07 Assets & Inventory", "Asset register (as of …).xlsx", "Every device ever installed"],
        ["08 Finance & Accounting", "YYYY/General ledger, Trial balance, Invoices (PDF), AR aging, Bank statements, "
                                    "Budget, Financial statements", "Double-entry books"],
        ["08 Finance & Accounting", "YYYY/Expense claims/Expense claims YYYY-MM.xlsx", "Staff expense claims (in GL 6200)"],
        ["09 HR & Workforce", "Employee roster, YYYY/Payroll register, Timesheets, SOC shift roster, Joiners and leavers, "
                              "Job descriptions", "People records (fictional)"],
        ["09 HR & Workforce", "Employee Files/<id name>/01 Contract & Onboarding, 02 Pay, 03 Performance, 04 Training, "
                              "05 Expenses, 06 Work Log, 07 Exit", "One file per person: offer, contract, IT onboarding, "
                              "pay statement, review, training, expenses, own alarms/tickets/work orders, exit"],
        ["11 Policies & SOPs", "IT SOPs/*.docx", "Step-by-step IT procedures (onboarding, patching, backups, P1 bridge…)"],
        ["12 Emails & Meetings", "YYYY/Zoom Meetings/<date title>/ (.vtt transcript, chat .txt, participants .csv)",
         "Weekly operations meetings, quarterly client service reviews"],
        ["12 Emails & Meetings", "YYYY/Google Meet/<date title>/ (transcript .docx, attendance .csv)",
         "P1 incident bridges, change advisory board"],
        ["10 Management Reports", "YYYY/Client service reports, KPI dashboards, Board packs", "Monthly and quarterly"],
        ["11 Policies & SOPs", "*.docx", "Policies and procedures"],
        ["12 Emails & Meetings", "YYYY/Email exports/*.zip (.eml), Meeting minutes/*.docx", "Correspondence"],
    ]
    wb = new_book()
    add_sheet(wb, "Datasets", ["Folder", "Files", "Contents"], dd, {"Folder": 30, "Files": 70, "Contents": 50})
    add_sheet(wb, "Chart of accounts", ["Account", "Name", "Type"], [[a, *v] for a, v in sorted(ACCOUNTS.items())],
              {"Name": 34})
    add_sheet(wb, "Notice", ["Notice"], [[BANNER], ["Domains use the reserved .example TLD; IBANs and serials are random."]],
              {"Notice": 100})
    save_book(sink, f"{ROOT}/Data dictionary.xlsx", wb, "Data dictionary", BANNER)


def build(sink, st, max_months=24, scale=1.0, detail=1.0):
    today = dt.date.today()
    last_full = ym_add((today.year, today.month), -1)
    done = tuple(st["done"]) if st.get("done") else None
    nxt = ym_add(done, 1) if done else START
    if nxt > last_full:
        return 0
    end = min(last_full, ym_add(nxt, max_months - 1))
    w = World(last_full, scale)
    ledger = st.setdefault("ledger", {})
    if not st.get("static"):
        static_docs(w, sink)
        st["static"] = True
    n = 0
    ym = nxt
    while ym <= end and not sink.full():
        Month(w, ym, ledger, sink, detail).run()
        st["done"] = list(ym)
        n += 1
        ym = ym_add(ym, 1)
    if n:
        last = tuple(st["done"])
        snapshots(w, last, sink)
        if last[1] != 12:  # employee files for the current year so far (rewritten in place next run)
            S.person_year(w, ledger, last[0], sink, sys.modules[__name__], final=False)
    return n
